"""The written deck: glued connectors, fit, colours, BPMN reading, lanes, strokes."""

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

import mermaid2pptx as m2p
from mermaid2pptx.flowchart import layout, parse, to_bpmn
from mermaid2pptx.flowchart import render as draw
from mermaid2pptx.flowchart.shapes import SHAPES
from mermaid2pptx.slide.deck import EMU_PER_PT
from ooxml import drawn_path, xfrm
from samples import LANES_SRC, diagram, examples


@pytest.mark.parametrize("path, direction, render", list(examples()))
def test_deck_connectors_are_glued_on_their_sites(path, direction, render, tmp_path):
    d = diagram(path, direction, render)
    lay = layout(d)
    prs = draw(d, lay, m2p.Options(render=render))
    out = tmp_path / "deck.pptx"
    prs.save(out)
    prs = Presentation(out)
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    assert len(slide.placeholders) == 0
    (grp,) = slide.shapes
    assert grp.shape_type == 6                          # MSO_SHAPE_TYPE.GROUP: one object to copy
    shapes = {s.shape_id: s for s in grp.shapes}
    nodes = {s.name[5:]: s for s in grp.shapes if s.name.startswith("node ")}
    assert set(nodes) == set(d.nodes)
    n_glued = 0
    for s in grp.shapes:
        if s._element.tag != qn("p:cxnSp"):
            continue
        el = s._element
        geom = el.find(f".//{qn('a:prstGeom')}")
        adj = [int(g.get("fmla").split()[1]) for g in geom.iter(qn("a:gd"))]
        x, y, w, h, rot, fv, fh = xfrm(el)
        drawn = drawn_path(geom.get("prst"), adj, x, y, w, h, rot, fv, fh)
        c_nv = el.find(f".//{qn('p:cNvCxnSpPr')}")
        for tag, point in (("a:stCxn", drawn[0]), ("a:endCxn", drawn[-1])):
            cxn = c_nv.find(qn(tag))
            assert cxn is not None, (s.name, tag)
            target = shapes[int(cxn.get("id"))]
            shape_name = d.nodes[target.name[5:]].shape
            bx, by, bw, bh, *_ = xfrm(target._element)
            sites = {idx: (bx + fx * bw, by + fy * bh) for idx, fx, fy in
                     ((i, fx(bw, bh) if callable(fx) else fx, fy)
                      for i, fx, fy in SHAPES[shape_name][1].values())}
            idx = int(cxn.get("idx"))
            assert idx in sites, (s.name, idx)
            # within 2pt: a jog shorter than that is snapped away rather than drawn
            assert point == pytest.approx(sites[idx], abs=2 * EMU_PER_PT), (s.name, tag)
            n_glued += 1
    visible = sum(1 for e in d.edges if e.line != "invisible")
    assert n_glued == 2 * visible


def test_fit_and_no_fit(tmp_path):
    src = "graph LR\n" + "\n".join(f"n{i}[A fairly long node label {i}] --> n{i + 1}" for i in range(14))
    prs, _, lay = m2p.convert(src)
    assert prs.slide_width == 12192000                  # 16:9, diagram shrunk to fit
    (grp,) = prs.slides[0].shapes
    assert grp.width <= prs.slide_width and grp.left >= 0
    prs, _, lay = m2p.convert(src, m2p.Options(fit=False, group=False))
    assert prs.slide_width > 12192000                   # slide grown to the natural size
    assert len(prs.slides[0].shapes) > 1


def test_theme_colors_and_no_theme_style(tmp_path):
    prs, _, _ = m2p.convert("graph TD\nA-->B", m2p.Options(color="theme", font="Arial"))
    xml = prs.slides[0].shapes._spTree.xml
    assert "schemeClr" in xml and 'typeface="Arial"' in xml
    assert "<p:style>" not in xml                        # no inherited theme shadow or fill


def _fills(prs):
    (grp,) = prs.slides[0].shapes
    return {s.name: str(s.fill.fore_color.rgb) for s in grp.shapes if s.name.startswith("node ")}


def test_colors():
    src = "graph TD\nA-->B\nstyle B fill:#FFEE00"
    prs, _, _ = m2p.convert(src, m2p.Options(render="mermaid"))
    assert _fills(prs) == {"node A": "5236AB", "node B": "FFEE00"}     # purple by default
    (grp,) = prs.slides[0].shapes
    text = {s.name: str(s.text_frame.paragraphs[0].runs[0].font.color.rgb) for s in grp.shapes
            if s.name.startswith("node ")}
    assert text == {"node A": "FFFFFF", "node B": "0F172A"}             # readable on either fill
    prs, _, _ = m2p.convert(src, m2p.Options(render="mermaid", color="#00A0B0"))
    assert _fills(prs)["node A"] == "00A0B0"
    with pytest.raises(ValueError, match="unknown colour"):
        m2p.palette("chartreuse-ish")


def test_bpmn_reading():
    src = """graph LR
    S([Début]) --> A[Saisie] --> G{OK ?}
    G -->|oui| P[[Sous-process]] --> DB[(Base)]
    G -->|non| A
    P --> M((Attente)) --> Z((Fin))
    """
    d = to_bpmn(parse(src))
    assert {n: v.shape for n, v in d.nodes.items()} == {
        "S": "start", "A": "task", "G": "gateway", "P": "call", "DB": "cylinder",
        "M": "intermediate", "Z": "end"}                  # explicit events: nothing added
    d = to_bpmn(parse("graph TD\nA[a] --> B{b}\nB --> C[c]\nB --> D[d]\nD --> B"))
    assert [n for n in d.nodes if d.nodes[n].shape == "start"] == ["start_A"]
    assert (d.edges[-2].src, d.edges[-2].dst) == ("start_A", "A")
    assert [n for n in d.nodes if d.nodes[n].shape == "end"] == ["end_C"]
    d = to_bpmn(parse("graph TD\nA --> B"), events=False)
    assert list(d.nodes) == ["A", "B"]


def test_bpmn_gateway_text_beside_the_symbol():
    prs, d, lay = m2p.convert("graph TD\nA --> G{Escalade nécessaire ?}\nG --> B")
    (grp,) = prs.slides[0].shapes
    shapes = {s.name: s for s in grp.shapes}
    diamond, text = shapes["node G"], shapes["text G"]
    assert diamond.width == diamond.height and diamond.text_frame.text == ""
    assert "Escalade" in text.text_frame.text
    assert text.left + text.width <= diamond.left                      # on its left in TD


def test_lanes_rendered_with_icons(tmp_path):
    prs, d, _ = m2p.convert(LANES_SRC)
    assert d.nodes["A"].text == "Demande" and d.nodes["A"].icon == "user"
    assert d.subgraphs["U"].title == "Client" and d.subgraphs["U"].icon == "user"
    (grp,) = prs.slides[0].shapes
    names = [s.name for s in grp.shapes]
    assert names[:2] == ["lane U", "lane header U"]      # lanes behind everything
    assert names.count("icon A") == 2 and names.count("icon U") == 2
    assert "👤" not in "".join(s.text_frame.text for s in grp.shapes if s.has_text_frame)
    d = to_bpmn(parse("graph TD\nsubgraph X\nA-->B\nend\nC-->A"))
    assert d.lanes == []                                 # C is outside: groups, not lanes
    d = to_bpmn(parse(LANES_SRC), lanes=False)
    assert d.lanes == []


def test_stroke_styles():
    src = ("graph TD\nA((a)):::s --> B:::e\nclassDef s stroke-dasharray: 0,stroke-width:2px\n"
           "classDef e stroke-dasharray: 5 5,stroke-width:3px")
    prs, _, _ = m2p.convert(src, m2p.Options(render="mermaid"))
    (grp,) = prs.slides[0].shapes
    nodes = {s.name: s for s in grp.shapes}
    assert nodes["node A"].line.dash_style is None and nodes["node A"].line.width == Pt(1.5)
    assert nodes["node B"].line.dash_style is not None and nodes["node B"].line.width == Pt(2.25)
