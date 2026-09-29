"""mermaid2pptx: parser, connector geometry, layout invariants and the written deck."""
import math
import random
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

import mermaid2pptx as m2p

EXAMPLES = sorted((Path(__file__).parent.parent / "examples").glob("*.mmd"))


# -- parser ----------------------------------------------------------------------------------

def test_shapes_and_labels():
    d = m2p.parse("""flowchart LR
        a[rect] --> b(round) --> c([stadium]) --> d[[sub]] --> e[(db)] --> f((circle))
        f --> g{diamond} --> h{{hex}} --> i[/para/] --> j[\\alt\\] --> k[/trap\\] --> l[\\trap alt/]
        l --> m>flag] --> n(((double))) --> o["quoted [brackets] text"]""")
    assert d.direction == "LR"
    shapes = {n.id: n.shape for n in d.nodes.values()}
    assert shapes == {"a": "rect", "b": "round", "c": "stadium", "d": "subroutine", "e": "cylinder",
                      "f": "circle", "g": "diamond", "h": "hexagon", "i": "parallelogram",
                      "j": "parallelogram_alt", "k": "trapezoid", "l": "trapezoid_alt", "m": "flag",
                      "n": "doublecircle", "o": "rect"}
    assert d.nodes["o"].text == "quoted [brackets] text"
    assert len(d.edges) == 14


def test_edge_kinds():
    d = m2p.parse("""graph TD
        A --> B
        A --- C
        A -.-> D
        A ==> E
        A ~~~ F
        A <--> G
        A --o H
        A --x I
        A ---> J
        A -- yes --> K
        A -->|no| L
        A -. maybe .-> M
        A == sure ==> N""")
    e = {x.dst: x for x in d.edges}
    assert (e["B"].line, e["B"].end, e["B"].minlen) == ("solid", "arrow", 1)
    assert (e["C"].end, e["C"].minlen) == ("none", 1)
    assert e["D"].line == "dotted" and e["E"].line == "thick" and e["F"].line == "invisible"
    assert (e["G"].start, e["G"].end) == ("arrow", "arrow")
    assert e["H"].end == "circle" and e["I"].end == "cross"
    assert e["J"].minlen == 2
    assert [e[k].text for k in "KLMN"] == ["yes", "no", "maybe", "sure"]
    assert e["M"].line == "dotted" and e["N"].line == "thick"


def test_chains_groups_and_statements():
    d = m2p.parse("""---
title: ignored
---
```mermaid
flowchart TB
    %% a comment
    A & B --> C & D; C --> E
```""")
    assert [(e.src, e.dst) for e in d.edges] == [("A", "C"), ("A", "D"), ("B", "C"), ("B", "D"), ("C", "E")]


def test_subgraphs_and_styles():
    d = m2p.parse("""flowchart TB
        subgraph outer [Outer box]
            A --> B
            subgraph inner
                C
            end
        end
        B --> C --> D:::hot
        classDef hot fill:#f96,stroke:#333,stroke-width:3px
        classDef default fill:#eee
        class A hot
        style B fill:#00ff00,color:white
        style outer fill:#fafafa
        linkStyle 0 stroke:#ff0000,stroke-width:2px""")
    assert d.path("A") == ("outer",) and d.path("C") == ("outer", "inner") and d.path("D") == ()
    assert d.subgraphs["outer"].title == "Outer box" and d.subgraphs["outer"].style["fill"] == "#fafafa"
    assert d.node_style(d.nodes["D"])["fill"] == "#f96"
    assert d.node_style(d.nodes["A"])["fill"] == "#f96"
    assert d.node_style(d.nodes["C"])["fill"] == "#eee"
    assert m2p.parse_color(d.node_style(d.nodes["B"])["color"]) == "FFFFFF"
    assert d.edges[0].style["stroke"] == "#ff0000"


def test_edge_to_subgraph_attaches_to_a_member():
    d = m2p.parse("""flowchart LR
        X --> grp
        subgraph grp
            A --> B
        end""")
    assert "grp" not in d.nodes
    assert (d.edges[0].src, d.edges[0].dst) == ("X", "A")
    assert d.warnings


def test_text_cleanup_and_colors():
    assert m2p.clean_text('"a<br>b<br/>c #quot;q#quot; &amp; #35;"') == 'a\nb\nc "q" & #'
    assert m2p.clean_text('"`**bold** md`"') == "bold md"
    assert m2p.parse_color("#abc") == "AABBCC"
    assert m2p.parse_color("rgb(255, 0, 16)") == "FF0010"
    assert m2p.parse_color("nonsense") is None


@pytest.mark.parametrize("src, msg", [
    ("sequenceDiagram\n A->>B: hi", "only 'flowchart'"),
    ("flowchart TD\n subgraph x\n A", "missing its 'end'"),
    ("flowchart TD\n A[oops --> B", "unclosed shape"),
    ("", "empty"),
])
def test_errors(src, msg):
    with pytest.raises(m2p.MermaidError, match=msg):
        m2p.parse(src)


def test_markdown_blocks():
    md = "text\n```mermaid\ngraph TD\nA-->B\n```\nmore\n```mermaid\ngraph LR\nC-->D\n```\n"
    blocks = m2p.extract_blocks(md)
    assert len(blocks) == 2 and "C-->D" in blocks[1]
    assert m2p.extract_blocks("graph TD\nA-->B") == ["graph TD\nA-->B"]


# -- connector geometry: an independent reading of the OOXML ---------------------------------

def drawn_path(prst, adj, x, y, w, h, rot=0, flip_v=False, flip_h=False):
    """The polyline PowerPoint draws for a connector: preset path, then flips, then rotation
    about the box centre (DrawingML semantics, written independently of the converter)."""
    a = [v / 100000 for v in adj]
    local = {
        "straightConnector1": [(0, 0), (w, h)],
        "bentConnector2": [(0, 0), (w, 0), (w, h)],
        "bentConnector3": [(0, 0), (w * a[0], 0), (w * a[0], h), (w, h)] if a else None,
        "bentConnector4": [(0, 0), (w * a[0], 0), (w * a[0], h * a[1]), (w, h * a[1]), (w, h)] if len(a) > 1 else None,
        "bentConnector5": [(0, 0), (w * a[0], 0), (w * a[0], h * a[1]), (w * a[2], h * a[1]),
                           (w * a[2], h), (w, h)] if len(a) > 2 else None,
    }[prst]
    t = math.radians(rot / 60000)
    out = []
    for px, py in local:
        if flip_h:
            px = w - px
        if flip_v:
            py = h - py
        dx, dy = px - w / 2, py - h / 2
        out.append((x + w / 2 + dx * math.cos(t) - dy * math.sin(t),
                    y + h / 2 + dx * math.sin(t) + dy * math.cos(t)))
    return out


def _random_route(rng, n):
    pts, (x, y) = [(0.0, 0.0)], (0.0, 0.0)
    axis = rng.choice((0, 1))
    for _ in range(n):
        step = rng.choice((-1, 1)) * rng.uniform(5, 100)
        x, y = (x + step, y) if axis == 0 else (x, y + step)
        pts.append((x, y))
        axis = 1 - axis
    return pts


@pytest.mark.parametrize("n", [1, 2, 3, 4, 5])
def test_connector_geometry_draws_the_route(n):
    rng = random.Random(n)
    for _ in range(300):
        pts = _random_route(rng, n)
        g = m2p.connector_geometry(pts)
        drawn = m2p.simplify(drawn_path(g["prst"], g["adj"], g["x"], g["y"], g["w"], g["h"], g["rot"], g["flipV"]))
        assert len(drawn) == len(pts)
        for p, q in zip(drawn, pts):     # 0.1pt: the minimum box extent of a bent connector
            assert p == pytest.approx(q, abs=0.11), (pts, g)


def test_connector_geometry_degenerate_width():
    """A back edge between two sites at the same x: the box gets a 0.1pt minimum width."""
    pts = [(100, 200), (130, 200), (130, 50), (100, 50)]
    g = m2p.connector_geometry(pts)
    drawn = drawn_path(g["prst"], g["adj"], g["x"], g["y"], g["w"], g["h"], g["rot"], g["flipV"])
    assert drawn[0] == pytest.approx(pts[0], abs=0.01)
    assert drawn[-1] == pytest.approx(pts[-1], abs=0.11)
    assert max(p[0] for p in drawn) == pytest.approx(130, abs=0.01)


def test_simplify():
    assert m2p.simplify([(0, 0), (0, 0), (0, 5), (0, 10), (3, 10)]) == [(0, 0), (0, 10), (3, 10)]


# -- layout invariants on every example, in every direction ----------------------------------

def _examples():
    for path in EXAMPLES:
        for direction in ("TB", "BT", "LR", "RL"):
            for render in ("bpmn", "mermaid"):
                yield pytest.param(path, direction, render, id=f"{path.stem}-{direction}-{render}")


def _diagram(path, direction, render):
    d = m2p.parse(path.read_text())
    d.direction = direction
    return m2p.to_bpmn(d) if render == "bpmn" else d


@pytest.mark.parametrize("path, direction, render", list(_examples()))
def test_layout_invariants(path, direction, render):
    d = _diagram(path, direction, render)
    lay = m2p.layout(d)
    boxes = list(lay.boxes.items())
    for i, (a, ba) in enumerate(boxes):                 # nodes never overlap
        for b, bb in boxes[i + 1:]:
            assert ba.x + ba.w <= bb.x + 0.01 or bb.x + bb.w <= ba.x + 0.01 or \
                ba.y + ba.h <= bb.y + 0.01 or bb.y + bb.h <= ba.y + 0.01, (a, b)
    for r in lay.routes:
        pts = r.points
        for p, q in zip(pts, pts[1:]):                  # orthogonal
            assert abs(p[0] - q[0]) < 0.01 or abs(p[1] - q[1]) < 0.01, (r.edge, pts)
        if r.edge.src == r.edge.dst:
            continue
        segments = list(zip(pts, pts[1:]))
        for n, b in lay.boxes.items():                  # never through a node
            for k, (p, q) in enumerate(segments):
                # the end segments may cut the bounding box corner of their own node to reach
                # an inset site (parallelogram, trapezoid)
                if (k == 0 and n == r.edge.src) or (k == len(segments) - 1 and n == r.edge.dst):
                    continue
                x0, x1, y0, y1 = min(p[0], q[0]), max(p[0], q[0]), min(p[1], q[1]), max(p[1], q[1])
                inside = x1 > b.x + 0.5 and x0 < b.x + b.w - 0.5 and y1 > b.y + 0.5 and y0 < b.y + b.h - 0.5
                assert not inside, (r.edge.src, r.edge.dst, "crosses", n)


# -- the written deck ------------------------------------------------------------------------

def _xfrm(el):
    x = el.find(f".//{qn('a:xfrm')}")
    off, ext = x.find(qn("a:off")), x.find(qn("a:ext"))
    return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")),
            int(x.get("rot", 0)), x.get("flipV") == "1", x.get("flipH") == "1")


@pytest.mark.parametrize("path, direction, render", list(_examples()))
def test_deck_connectors_are_glued_on_their_sites(path, direction, render, tmp_path):
    d = _diagram(path, direction, render)
    lay = m2p.layout(d)
    prs = m2p.render(d, lay, m2p.Options(render=render))
    out = tmp_path / "deck.pptx"
    prs.save(out)
    prs = Presentation(out)
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    assert len(slide.placeholders) == 0
    (grp,) = slide.shapes
    assert grp.shape_type == 6                          # MSO_SHAPE_TYPE.GROUP: one object to copy
    shapes = {s.shape_id: s for s in grp.shapes}
    nodes = {s.name[5:]: s for s in grp.shapes if s.name.startswith("node ")}
    assert set(nodes) == set(d.nodes)
    n_glued = 0
    for s in grp.shapes:
        if s._element.tag != qn("p:cxnSp"):
            continue
        el = s._element
        geom = el.find(f".//{qn('a:prstGeom')}")
        adj = [int(g.get("fmla").split()[1]) for g in geom.iter(qn("a:gd"))]
        x, y, w, h, rot, fv, fh = _xfrm(el)
        drawn = drawn_path(geom.get("prst"), adj, x, y, w, h, rot, fv, fh)
        c_nv = el.find(f".//{qn('p:cNvCxnSpPr')}")
        for tag, point in (("a:stCxn", drawn[0]), ("a:endCxn", drawn[-1])):
            cxn = c_nv.find(qn(tag))
            assert cxn is not None, (s.name, tag)
            target = shapes[int(cxn.get("id"))]
            shape_name = d.nodes[target.name[5:]].shape
            bx, by, bw, bh, *_ = _xfrm(target._element)
            sites = {idx: (bx + fx * bw, by + fy * bh) for idx, fx, fy in
                     ((i, fx(bw, bh) if callable(fx) else fx, fy)
                      for i, fx, fy in m2p.SHAPES[shape_name][1].values())}
            idx = int(cxn.get("idx"))
            assert idx in sites, (s.name, idx)
            # within 2pt: a jog shorter than that is snapped away rather than drawn
            assert point == pytest.approx(sites[idx], abs=2 * m2p.EMU_PER_PT), (s.name, tag)
            n_glued += 1
    visible = sum(1 for e in d.edges if e.line != "invisible")
    assert n_glued == 2 * visible


def test_fit_and_no_fit(tmp_path):
    src = "graph LR\n" + "\n".join(f"n{i}[A fairly long node label {i}] --> n{i + 1}" for i in range(14))
    prs, _, lay = m2p.convert(src)
    assert prs.slide_width == 12192000                  # 16:9, diagram shrunk to fit
    (grp,) = prs.slides[0].shapes
    assert grp.width <= prs.slide_width and grp.left >= 0
    prs, _, lay = m2p.convert(src, m2p.Options(fit=False, group=False))
    assert prs.slide_width > 12192000                   # slide grown to the natural size
    assert len(prs.slides[0].shapes) > 1


def test_theme_colors_and_no_theme_style(tmp_path):
    prs, _, _ = m2p.convert("graph TD\nA-->B", m2p.Options(color="theme", font="Arial"))
    xml = prs.slides[0].shapes._spTree.xml
    assert "schemeClr" in xml and 'typeface="Arial"' in xml
    assert "<p:style>" not in xml                        # no inherited theme shadow or fill


def _fills(prs):
    (grp,) = prs.slides[0].shapes
    return {s.name: str(s.fill.fore_color.rgb) for s in grp.shapes if s.name.startswith("node ")}


def test_colors():
    src = "graph TD\nA-->B\nstyle B fill:#FFEE00"
    prs, _, _ = m2p.convert(src, m2p.Options(render="mermaid"))
    assert _fills(prs) == {"node A": "5236AB", "node B": "FFEE00"}     # purple by default
    (grp,) = prs.slides[0].shapes
    text = {s.name: str(s.text_frame.paragraphs[0].runs[0].font.color.rgb) for s in grp.shapes
            if s.name.startswith("node ")}
    assert text == {"node A": "FFFFFF", "node B": "0F172A"}             # readable on either fill
    prs, _, _ = m2p.convert(src, m2p.Options(render="mermaid", color="#00A0B0"))
    assert _fills(prs)["node A"] == "00A0B0"
    with pytest.raises(ValueError, match="unknown colour"):
        m2p.palette("chartreuse-ish")


def test_bpmn_reading():
    src = """graph LR
    S([Début]) --> A[Saisie] --> G{OK ?}
    G -->|oui| P[[Sous-process]] --> DB[(Base)]
    G -->|non| A
    P --> M((Attente)) --> Z((Fin))
    """
    d = m2p.to_bpmn(m2p.parse(src))
    assert {n: v.shape for n, v in d.nodes.items()} == {
        "S": "start", "A": "task", "G": "gateway", "P": "call", "DB": "cylinder",
        "M": "intermediate", "Z": "end"}                  # explicit events: nothing added
    d = m2p.to_bpmn(m2p.parse("graph TD\nA[a] --> B{b}\nB --> C[c]\nB --> D[d]\nD --> B"))
    assert [n for n in d.nodes if d.nodes[n].shape == "start"] == ["start_A"]
    assert (d.edges[-2].src, d.edges[-2].dst) == ("start_A", "A")
    assert [n for n in d.nodes if d.nodes[n].shape == "end"] == ["end_C"]
    d = m2p.to_bpmn(m2p.parse("graph TD\nA --> B"), events=False)
    assert list(d.nodes) == ["A", "B"]


def test_bpmn_gateway_text_beside_the_symbol():
    prs, d, lay = m2p.convert("graph TD\nA --> G{Escalade nécessaire ?}\nG --> B")
    (grp,) = prs.slides[0].shapes
    shapes = {s.name: s for s in grp.shapes}
    diamond, text = shapes["node G"], shapes["text G"]
    assert diamond.width == diamond.height and diamond.text_frame.text == ""
    assert "Escalade" in text.text_frame.text
    assert text.left + text.width <= diamond.left                      # on its left in TD


def test_author(tmp_path, monkeypatch):
    monkeypatch.delenv(m2p.AUTHOR_ENV, raising=False)
    prs, _, _ = m2p.convert("graph TD\nA-->B")
    out = tmp_path / "a.pptx"
    prs.save(out)
    cp = Presentation(out).core_properties
    assert cp.author == cp.last_modified_by == ""                    # not python-pptx's own
    assert cp.comments == "" and cp.revision == 1
    monkeypatch.setenv(m2p.AUTHOR_ENV, "From Env")
    prs, _, _ = m2p.convert("graph TD\nA-->B")
    assert prs.core_properties.author == "From Env"
    prs, _, _ = m2p.convert("graph TD\nA-->B", m2p.Options(author="Someone Else"))
    assert prs.core_properties.author == "Someone Else"


def test_cli(tmp_path, capsys):
    md = tmp_path / "doc.md"
    md.write_text("```mermaid\ngraph TD\nA-->B\n```\n```mermaid\ngraph LR\nX-->Y-->Z\n```\n")
    assert m2p.main([str(md), "--block", "2"]) == 0
    assert (tmp_path / "doc.pptx").exists()
    assert "5 nodes, 4 edges" in capsys.readouterr().out            # + start and end events
    assert m2p.main([str(md), "--block", "2", "--render", "mermaid", "--color", "slate"]) == 0
    assert "3 nodes, 2 edges" in capsys.readouterr().out
    with pytest.raises(SystemExit):
        m2p.main([str(md), "--color", "nope"])
    bad = tmp_path / "bad.mmd"
    bad.write_text("pie\n")
    assert m2p.main([str(bad), "-o", str(tmp_path / "x.pptx")]) == 1
    with pytest.raises(SystemExit):
        m2p.main([str(md), "--block", "3"])
    with pytest.raises(SystemExit):
        m2p.main(["--version"])
    assert m2p.__version__ in capsys.readouterr().out


LANES_SRC = """graph TD
    subgraph U["👤 Client"]
        S(( )) --> A["👤 Demande"]
    end
    subgraph T[Outil]
        B{Type ?}
        subgraph N[Nested]
            C[Ticket]
        end
    end
    subgraph P[Support]
        D[Traitement] --> E(( ))
        D --> D
    end
    A --> B
    B -->|incident| C
    B -->|demande| D
    C --> D
    D --> A
"""


@pytest.mark.parametrize("direction", ["TB", "BT", "LR", "RL"])
def test_lanes(direction):
    d = m2p.to_bpmn(m2p.parse(LANES_SRC))
    d.direction = direction
    assert d.lanes == ["U", "T", "P"]
    lay = m2p.layout(d)
    vertical = direction in ("TB", "BT")
    lanes = {sid: body for sid, body, _, _ in lay.lanes}
    spans = [(b.x, b.x + b.w) if vertical else (b.y, b.y + b.h) for b in lanes.values()]
    for (_, a1), (b0, _) in zip(spans, spans[1:]):     # adjacent, in declaration order
        assert b0 == pytest.approx(a1)
    for n, box in lay.boxes.items():                    # each node inside its own lane
        body = lanes[d.path(n)[0]]
        assert body.x - 0.01 <= box.x and box.x + box.w <= body.x + body.w + 0.01, n
        assert body.y - 0.01 <= box.y and box.y + box.h <= body.y + body.h + 0.01, n
    headers = [h for _, _, h, _ in lay.lanes]
    assert all((h.y + h.h if vertical else h.x + h.w) == pytest.approx(
        lanes["U"].y if vertical else lanes["U"].x) for h in headers)   # headers before the flow
    assert [c[0] for c in lay.clusters] == ["N"]        # a nested subgraph stays a group


def test_lanes_rendered_with_icons(tmp_path):
    prs, d, _ = m2p.convert(LANES_SRC)
    assert d.nodes["A"].text == "Demande" and d.nodes["A"].icon == "user"
    assert d.subgraphs["U"].title == "Client" and d.subgraphs["U"].icon == "user"
    (grp,) = prs.slides[0].shapes
    names = [s.name for s in grp.shapes]
    assert names[:2] == ["lane U", "lane header U"]      # lanes behind everything
    assert names.count("icon A") == 2 and names.count("icon U") == 2
    assert "👤" not in "".join(s.text_frame.text for s in grp.shapes if s.has_text_frame)
    d = m2p.to_bpmn(m2p.parse("graph TD\nsubgraph X\nA-->B\nend\nC-->A"))
    assert d.lanes == []                                 # C is outside: groups, not lanes
    d = m2p.to_bpmn(m2p.parse(LANES_SRC), lanes=False)
    assert d.lanes == []


def test_stroke_styles():
    src = ("graph TD\nA((a)):::s --> B:::e\nclassDef s stroke-dasharray: 0,stroke-width:2px\n"
           "classDef e stroke-dasharray: 5 5,stroke-width:3px")
    prs, _, _ = m2p.convert(src, m2p.Options(render="mermaid"))
    (grp,) = prs.slides[0].shapes
    nodes = {s.name: s for s in grp.shapes}
    assert nodes["node A"].line.dash_style is None and nodes["node A"].line.width == m2p.Pt(1.5)
    assert nodes["node B"].line.dash_style is not None and nodes["node B"].line.width == m2p.Pt(2.25)
