"""--template: the deck takes a .pptx or .potx's theme and slide size, so --color theme shows its colours."""
import zipfile

import pytest
from pptx import Presentation
from pptx.util import Emu

import mermaid2pptx as m2p
from mermaid2pptx import cli

ACCENT = "C0FFEE"


def make_template(tmp_path, suffix=".pptx"):
    """A 4:3 deck with accent 1 set to ACCENT, a title slide, and no layout without placeholders."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(6858000)
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Old slide"
    blank = prs.slide_layouts[6]
    prs.slide_layouts.remove(blank)
    path = tmp_path / "t.pptx"
    prs.save(path)
    out = tmp_path / f"template{suffix}"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item)
            if item.filename.startswith("ppt/theme/theme"):
                data = data.replace(b'<a:accent1><a:srgbClr val="4F81BD"/>',
                                    f'<a:accent1><a:srgbClr val="{ACCENT}"/>'.encode())
            if item.filename == "[Content_Types].xml" and suffix == ".potx":
                data = data.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            dst.writestr(item, data)
    return out


def theme_xml(path):
    with zipfile.ZipFile(path) as z:
        return b"".join(z.read(n) for n in z.namelist() if n.startswith("ppt/theme/theme"))


@pytest.mark.parametrize("suffix", [".pptx", ".potx"])
def test_template_theme_and_size(tmp_path, suffix, monkeypatch):
    monkeypatch.delenv(m2p.TEMPLATE_ENV, raising=False)
    template = make_template(tmp_path, suffix)
    assert ACCENT.encode() in theme_xml(template)
    prs, _, _ = m2p.convert("graph TD\nA-->B", m2p.Options(color="theme", template=str(template)))
    out = tmp_path / "out.pptx"
    prs.save(out)
    deck = Presentation(out)
    assert ACCENT.encode() in theme_xml(out)
    assert (deck.slide_width, deck.slide_height) == (9144000, 6858000)       # the template's, not 16:9
    assert len(deck.slides) == 1                                             # its own slides are gone
    slide = deck.slides[0]
    assert len(slide.placeholders) == 0 and len(slide.shapes) == 1           # only the diagram's group
    with zipfile.ZipFile(out) as z:
        assert not any(b"Old slide" in z.read(n) for n in z.namelist() if n.startswith("ppt/slides/"))


def test_template_from_env_and_cli(tmp_path, capsys, monkeypatch):
    template = make_template(tmp_path, ".potx")
    mmd = tmp_path / "d.mmd"
    mmd.write_text("graph TD\nA-->B\n", encoding="utf-8")
    monkeypatch.setenv(m2p.TEMPLATE_ENV, str(template))
    assert cli.main([str(mmd), "--color", "theme"]) == 0                     # the Windows dialog's route
    assert ACCENT.encode() in theme_xml(tmp_path / "d.pptx")
    assert cli.main([str(mmd), "--template", ""]) == 0                       # "" = the built-in deck
    assert ACCENT.encode() not in theme_xml(tmp_path / "d.pptx")
    monkeypatch.delenv(m2p.TEMPLATE_ENV)
    assert cli.main([str(mmd), "--template", str(template)]) == 0
    assert ACCENT.encode() in theme_xml(tmp_path / "d.pptx")
    capsys.readouterr()
    for bad in (tmp_path / "missing.pptx", mmd):
        assert cli.main([str(mmd), "--template", str(bad)]) == 1
        assert capsys.readouterr().err.startswith(f"error: template {bad}: not a PowerPoint")
