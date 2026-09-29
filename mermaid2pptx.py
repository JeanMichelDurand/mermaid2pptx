#!/usr/bin/env python3
"""Convert a Mermaid flowchart into native, editable PowerPoint shapes on one blank slide.

    mermaid2pptx diagram.mmd -o diagram.pptx
    mermaid2pptx notes.md --block 2                    # 2nd ```mermaid block
    cat diagram.mmd | mermaid2pptx - -o diagram.pptx

The deck is python-pptx's built-in blank presentation, not a corporate template: open it,
click the diagram (a single group) and paste it into any slide. Nodes are autoshapes;
edges are connectors glued to the nodes' connection sites, so they follow a node you move.

Only `flowchart` / `graph` diagrams are supported. The layout is a small layered
(Sugiyama-style) engine written here, with orthogonal routing: every edge is a straight
line or an elbow connector whose segments are horizontal/vertical, with horizontal jogs
spread on separate tracks between two layers so they never sit on top of each other.
By default (--render bpmn) the flowchart is read as a simplified BPMN process: decisions
become gateways, circles and stadiums events, other boxes tasks, with missing start/end
events added, top-level subgraphs drawn as swim lanes and a leading 👤 as a user icon
(to_bpmn). Boxes are purple unless --color says otherwise.
See README.md for the syntax covered and the known limits.
"""
from __future__ import annotations

import argparse
import html
import os
import re
import sys
from dataclasses import dataclass, field, replace
from datetime import datetime, timezone
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

__version__ = "1.0.1"

EMU_PER_PT = 12700


# --------------------------------------------------------------------------------------------
# Model
# --------------------------------------------------------------------------------------------

@dataclass
class Node:
    id: str
    text: str
    shape: str = "rect"
    classes: list[str] = field(default_factory=list)
    style: dict[str, str] = field(default_factory=dict)
    cluster: str | None = None          # innermost subgraph id
    icon: str | None = None             # bpmn: "user" (a leading 👤 in the text)


@dataclass
class Edge:
    src: str
    dst: str
    text: str = ""
    line: str = "solid"                 # solid | dotted | thick | invisible
    start: str = "none"                 # marker at src: none | arrow | circle | cross
    end: str = "arrow"                  # marker at dst
    minlen: int = 1
    style: dict[str, str] = field(default_factory=dict)


@dataclass
class Subgraph:
    id: str
    title: str
    parent: str | None
    style: dict[str, str] = field(default_factory=dict)
    icon: str | None = None


@dataclass
class Diagram:
    direction: str = "TB"
    nodes: dict[str, Node] = field(default_factory=dict)
    edges: list[Edge] = field(default_factory=list)
    subgraphs: dict[str, Subgraph] = field(default_factory=dict)
    class_defs: dict[str, dict[str, str]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    lanes: list[str] = field(default_factory=list)      # bpmn: top-level subgraphs drawn as lanes

    def path(self, node_id: str) -> tuple[str, ...]:
        """Subgraph ids enclosing the node, outermost first."""
        out, sg = [], self.nodes[node_id].cluster
        while sg:
            out.append(sg)
            sg = self.subgraphs[sg].parent
        return tuple(reversed(out))

    def node_style(self, node: Node) -> dict[str, str]:
        style = dict(self.class_defs.get("default", {}))
        for cls in node.classes:
            style.update(self.class_defs.get(cls, {}))
        style.update(node.style)
        return style


class MermaidError(ValueError):
    pass


# --------------------------------------------------------------------------------------------
# Parser
# --------------------------------------------------------------------------------------------

DIRECTIONS = {"TB": "TB", "TD": "TB", "BT": "BT", "LR": "LR", "RL": "RL",
              "V": "TB", "^": "BT", ">": "LR", "<": "RL"}

_ID = re.compile(r"\w+(?:[-.]\w+)*")
# opener -> [(closer, shape)], longest openers first
_OPENERS = [
    ("(((", [(")))", "doublecircle")]),
    ("((", [("))", "circle")]),
    ("([", [("])", "stadium")]),
    ("[[", [("]]", "subroutine")]),
    ("[(", [(")]", "cylinder")]),
    ("{{", [("}}", "hexagon")]),
    ("[/", [("/]", "parallelogram"), ("\\]", "trapezoid")]),
    ("[\\", [("\\]", "parallelogram_alt"), ("/]", "trapezoid_alt")]),
    ("(", [(")", "round")]),
    ("[", [("]", "rect")]),
    ("{", [("}", "diamond")]),
    (">", [("]", "flag")]),
]
# `A -- text --> B`, `A -. text .-> B`, `A == text ==> B`
_EDGE_TEXT = re.compile(
    r"(?P<l><|[ox](?=[-=]))?(?P<open>--|-\.|==)(?![->.=]|[ox]\s)\s*(?P<text>\S.*?)\s*"
    r"(?P<close>-{2,}[>ox]?|\.+-[>ox]?|={2,}[>ox]?)(?=\s|$|[\w\"])")
# `-->`, `---`, `-.->`, `==>`, `~~~`, `<-->`, `--o`, `--x`, `--->` (longer = more ranks)
_EDGE = re.compile(r"(?P<l><|[ox](?=[-=]))?(?P<body>-{2,}|-\.+-|={2,}|~{3,})(?P<r>>|[ox](?=\s|$))?")
_PIPE = re.compile(r"\s*\|(?P<text>[^|]*)\|")
_MARKERS = {"": "none", None: "none", ">": "arrow", "<": "arrow", "o": "circle", "x": "cross"}

_NAMED_COLORS = {
    "white": "FFFFFF", "black": "000000", "red": "FF0000", "green": "008000", "blue": "0000FF",
    "yellow": "FFFF00", "orange": "FFA500", "purple": "800080", "gray": "808080", "grey": "808080",
    "lightgray": "D3D3D3", "lightgrey": "D3D3D3", "darkgray": "A9A9A9", "pink": "FFC0CB",
    "lightblue": "ADD8E6", "lightgreen": "90EE90", "navy": "000080", "teal": "008080",
}


def clean_text(text: str) -> str:
    """Mermaid label -> plain text with '\\n' line breaks."""
    t = text.strip()
    if len(t) >= 2 and t[0] == t[-1] == '"':
        t = t[1:-1]
    if len(t) >= 2 and t[0] == t[-1] == "`":            # markdown string
        t = t[1:-1].replace("**", "").replace("__", "")
    t = re.sub(r"<br\s*/?>", "\n", t, flags=re.I)
    t = re.sub(r"#(\w+);", lambda m: f"&#{m[1]};" if m[1].isdigit() else f"&{m[1]};", t)
    t = html.unescape(t)
    t = re.sub(r"</?[a-zA-Z][^>]*>", "", t)
    t = re.sub(r"\bfa[bsr]?:fa-[\w-]+\s*", "", t)         # font-awesome icons
    return "\n".join(line.strip() for line in t.split("\n")).strip()


def parse_color(value: str) -> str | None:
    v = value.strip().lower()
    if m := re.fullmatch(r"#([0-9a-f]{3})", v):
        return "".join(c * 2 for c in m[1]).upper()
    if m := re.fullmatch(r"#([0-9a-f]{6})([0-9a-f]{2})?", v):
        return m[1].upper()
    if m := re.fullmatch(r"rgba?\((\d+)\s*,\s*(\d+)\s*,\s*(\d+).*\)", v):
        return "".join(f"{min(int(x), 255):02X}" for x in m.groups())
    return _NAMED_COLORS.get(v)


def parse_style(spec: str) -> dict[str, str]:
    out = {}
    for part in re.split(r",(?![^(]*\))", spec):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = v.strip().rstrip(";")
    return out


def _split_statements(line: str) -> list[str]:
    """Split on ';' outside quotes and brackets."""
    out, cur, depth, quote = [], [], 0, False
    for ch in line:
        if ch == '"':
            quote = not quote
        elif not quote and ch in "([{":
            depth += 1
        elif not quote and ch in ")]}":
            depth = max(0, depth - 1)
        elif ch == ";" and not quote and depth == 0:
            out.append("".join(cur))
            cur = []
            continue
        cur.append(ch)
    out.append("".join(cur))
    return [s.strip() for s in out if s.strip()]


def extract_blocks(src: str) -> list[str]:
    """The ```mermaid blocks of a Markdown text, or the text itself if it has none."""
    blocks = re.findall(r"^\s*(?:```|~~~)\s*mermaid[^\n]*\n(.*?)^\s*(?:```|~~~)", src, re.M | re.S)
    return blocks or [src]


class _Parser:
    def __init__(self):
        self.d = Diagram()
        self.stack: list[str] = []
        self.explicit: set[str] = set()       # ids given a label or shape somewhere
        self.sg_counter = 0

    # -- statements -------------------------------------------------------------------------
    def parse(self, src: str) -> Diagram:
        lines = src.replace("\r\n", "\n").split("\n")
        lines = [ln for ln in lines if not ln.strip().startswith(("```", "~~~"))]
        while lines and not lines[0].strip():
            lines.pop(0)
        if lines and lines[0].strip() == "---":                  # YAML front matter
            end = next((i for i, ln in enumerate(lines[1:], 1) if ln.strip() == "---"), 0)
            lines = lines[end + 1:]
        stmts = [s for ln in lines if ln.strip() and not ln.strip().startswith("%%")
                 for s in _split_statements(ln.strip())]
        if not stmts:
            raise MermaidError("empty diagram")
        m = re.fullmatch(r"(flowchart|graph)(?:\s+(\S+))?", stmts[0], re.I)
        if not m:
            kind = stmts[0].split()[0]
            raise MermaidError(f"only 'flowchart'/'graph' diagrams are supported, not {kind!r}")
        if m[2]:
            if m[2].upper() not in DIRECTIONS:
                raise MermaidError(f"unknown direction {m[2]!r}")
            self.d.direction = DIRECTIONS[m[2].upper()]
        for st in stmts[1:]:
            try:
                self.statement(st)
            except MermaidError as exc:
                raise MermaidError(f"{exc} in: {st}") from None
        if self.stack:
            raise MermaidError(f"subgraph {self.stack[-1]!r} is missing its 'end'")
        self.resolve_subgraph_endpoints()
        return self.d

    def statement(self, st: str):
        word = st.split(None, 1)[0]
        low = word.lower()
        rest = st[len(word):].strip()
        if low == "subgraph":
            self.subgraph(rest)
        elif low == "end" and not rest:
            if not self.stack:
                raise MermaidError("'end' without subgraph")
            self.stack.pop()
        elif low == "direction" or low in ("click", "call", "href") or low.startswith(("acctitle", "accdescr")):
            return
        elif low == "classdef":
            names, _, spec = rest.partition(" ")
            for name in names.split(","):
                self.d.class_defs.setdefault(name.strip(), {}).update(parse_style(spec))
        elif low == "class":
            ids, _, cls = rest.rpartition(" ")
            for nid in ids.split(","):
                self.touch(nid.strip()).classes.append(cls.strip())
        elif low == "style":
            nid, _, spec = rest.partition(" ")
            target = self.d.subgraphs.get(nid) or self.touch(nid.strip())
            target.style.update(parse_style(spec))
        elif low == "linkstyle":
            which, _, spec = rest.partition(" ")
            idx = range(len(self.d.edges)) if which == "default" else \
                [int(i) for i in which.split(",") if i.strip().isdigit()]
            for i in idx:
                if i < len(self.d.edges):
                    self.d.edges[i].style.update(parse_style(spec))
        else:
            self.chain(st)

    def subgraph(self, rest: str):
        if m := re.fullmatch(r"(\w[\w.-]*)\s*\[(.*)\]", rest):
            sid, title = m[1], clean_text(m[2])
        elif rest.startswith('"'):
            self.sg_counter += 1
            sid, title = f"subGraph{self.sg_counter}", clean_text(rest)
        elif rest and " " not in rest:
            sid = title = rest
        else:
            self.sg_counter += 1
            sid, title = (rest if rest else f"subGraph{self.sg_counter}"), clean_text(rest)
        self.d.subgraphs[sid] = Subgraph(sid, title, self.stack[-1] if self.stack else None)
        self.stack.append(sid)

    # -- chains: A[x] --> B & C -->|y| D ----------------------------------------------------
    def chain(self, st: str):
        i = 0
        group, i = self.group(st, i)
        while True:
            i = _skip_ws(st, i)
            if i >= len(st):
                return
            edge, i = self.edge(st, i)
            i = _skip_ws(st, i)
            nxt, i = self.group(st, i)
            for a in group:
                for b in nxt:
                    self.d.edges.append(Edge(a, b, edge["text"], edge["line"], edge["start"],
                                             edge["end"], edge["minlen"]))
            group = nxt

    def group(self, st: str, i: int) -> tuple[list[str], int]:
        ids = []
        while True:
            nid, i = self.node(st, _skip_ws(st, i))
            ids.append(nid)
            j = _skip_ws(st, i)
            if j < len(st) and st[j] == "&":
                i = j + 1
                continue
            return ids, i

    def node(self, st: str, i: int) -> tuple[str, int]:
        m = _ID.match(st, i)
        if not m:
            raise MermaidError(f"expected a node id at {st[i:i + 20]!r}")
        nid, i = m[0], m.end()
        text = shape = None
        for opener, closers in _OPENERS:
            if not st.startswith(opener, i):
                continue
            j = i + len(opener)
            if j < len(st) and st[j] == '"':
                q = st.find('"', j + 1)
                if q < 0:
                    raise MermaidError("unterminated string")
                text, k = st[j + 1:q], q + 1
                hits = [(k, c, s) for c, s in closers if st.startswith(c, k)]
            else:
                hits = sorted((st.find(c, j), c, s) for c, s in closers if st.find(c, j) >= 0)
                if hits:
                    text = st[j:hits[0][0]]
            if not hits:
                raise MermaidError(f"unclosed shape {opener!r} for node {nid!r}")
            k, closer, shape = hits[0]
            i = k + len(closer)
            break
        classes = []
        while st.startswith(":::", i):
            m = _ID.match(st, i + 3)
            if not m:
                break
            classes.append(m[0])
            i = m.end()
        node = self.touch(nid)
        if text is not None:
            node.text = clean_text(text)
            self.explicit.add(nid)
        if shape:
            node.shape = shape
            self.explicit.add(nid)
        node.classes += classes
        return nid, i

    def edge(self, st: str, i: int) -> tuple[dict, int]:
        if m := _EDGE_TEXT.match(st, i):
            close = m["close"]
            head = close[-1] if close[-1] in ">ox" else ""
            body = close[:-1] if head else close
            text, kind = clean_text(m["text"]), m["open"]
            end = m.end()
        elif m := _EDGE.match(st, i):
            head, body, kind, text, end = m["r"] or "", m["body"], m["body"], "", m.end()
        else:
            raise MermaidError(f"expected an edge at {st[i:i + 20]!r}")
        if p := _PIPE.match(st, end):
            text, end = clean_text(p["text"]), p.end()
        left = m["l"] or ""
        if "." in kind:
            line, minlen = "dotted", body.count(".")
        elif kind.startswith("~"):
            line, minlen = "invisible", 1
        else:           # `-->`/`==>` is 1 rank, `---`/`===` too; each extra dash adds one
            line = "thick" if kind.startswith("=") else "solid"
            minlen = len(body) - (1 if head or (left and m.re is _EDGE) else 2)
        return dict(text=text, line=line, start=_MARKERS[left], end=_MARKERS[head],
                    minlen=max(1, minlen)), end

    # -- nodes ------------------------------------------------------------------------------
    def touch(self, nid: str) -> Node:
        node = self.d.nodes.get(nid)
        if node is None:
            node = self.d.nodes[nid] = Node(nid, nid)
        if self.stack and node.cluster is None and nid not in self.d.subgraphs:
            node.cluster = self.stack[-1]
        return node

    def resolve_subgraph_endpoints(self):
        """`A --> subgraphId`: mermaid draws to the cluster; we attach to its first/last node."""
        d = self.d
        for sid in d.subgraphs:
            if sid in d.nodes and sid not in self.explicit:
                del d.nodes[sid]

        def members(sid):
            subs = {s for s in d.subgraphs if s == sid or sid in _ancestors(d, s)}
            return [n for n, node in d.nodes.items() if node.cluster in subs]

        kept = []
        for e in d.edges:
            ok = True
            for attr, pick in (("src", -1), ("dst", 0)):
                sid = getattr(e, attr)
                if sid in d.subgraphs and sid not in d.nodes:
                    ms = members(sid)
                    if not ms:
                        d.warnings.append(f"edge to empty subgraph {sid!r} dropped")
                        ok = False
                        break
                    setattr(e, attr, ms[pick])
                    d.warnings.append(f"edge to subgraph {sid!r} attached to node {ms[pick]!r}")
            if ok:
                kept.append(e)
        d.edges = kept


def _ancestors(d: Diagram, sid: str) -> list[str]:
    out, p = [], d.subgraphs[sid].parent
    while p:
        out.append(p)
        p = d.subgraphs[p].parent
    return out


def _skip_ws(s: str, i: int) -> int:
    while i < len(s) and s[i].isspace():
        i += 1
    return i


def parse(src: str) -> Diagram:
    return _Parser().parse(src)


# --------------------------------------------------------------------------------------------
# BPMN reading
# --------------------------------------------------------------------------------------------

_DATA = {"cylinder", "parallelogram", "parallelogram_alt"}     # data store, data object: kept
_USER = re.compile(r"^\s*[\U0001F464\U0001F465]\uFE0F?\s*")    # 👤 👥


def _take_icon(text: str) -> tuple[str, str | None]:
    m = _USER.match(text)
    return (text[m.end():], "user") if m else (text, None)


def to_bpmn(d: Diagram, events: bool = True, lanes: bool = True) -> Diagram:
    """Read the flowchart as a simplified BPMN process (in place, and returned).

    `{decision}` -> exclusive gateway. `((circle))`, `(((double circle)))` and `([stadium])`
    -> events: start without incoming flow, end without outgoing flow (a double circle is
    always an end), intermediate otherwise (a stadium in mid-flow stays a task).
    `[[subroutine]]` -> call activity. Database and parallelogram shapes are kept. Every other
    shape -> task. With `events`, a process with no start (end) event gets one before each
    source (after each sink). A leading 👤 in a text becomes a user icon. With `lanes`, when
    every node sits in one of two or more top-level subgraphs, those are drawn as lanes.
    """
    for n in d.nodes.values():
        n.text, n.icon = _take_icon(n.text)
    for sg in d.subgraphs.values():
        sg.title, sg.icon = _take_icon(sg.title)
    tops = [s for s, sg in d.subgraphs.items() if sg.parent is None]
    if lanes and len(tops) >= 2 and all(d.path(n)[:1] for n in d.nodes):
        d.lanes = tops
    flows = [e for e in d.edges if e.line != "invisible" and e.src != e.dst]
    has_in, has_out = {e.dst for e in flows}, {e.src for e in flows}
    for n in d.nodes.values():
        sh, i, o = n.shape, n.id in has_in, n.id in has_out
        if sh == "diamond":
            n.shape = "gateway"
        elif sh == "doublecircle" or (sh in ("circle", "stadium") and i and not o):
            n.shape = "end"
        elif sh in ("circle", "stadium") and not i:
            n.shape = "start"
        elif sh == "circle":
            n.shape = "intermediate"
        elif sh == "subroutine":
            n.shape = "call"
        elif sh not in _DATA:
            n.shape = "task"
    if not events:
        return d
    kinds = {n.shape for n in d.nodes.values()}
    steps = [n for n in d.nodes.values() if n.shape in ("task", "call", "gateway")]
    sources = [] if "start" in kinds else [n for n in steps if n.id not in has_in and n.id in has_out]
    sinks = [] if "end" in kinds else [n for n in steps if n.id in has_in and n.id not in has_out]

    def fresh(base):
        while base in d.nodes or base in new:
            base += "_"
        return base

    new: dict[str, Node] = {}
    starts, ends, e_in, e_out = {}, {}, [], []
    for n in sources:
        k = fresh(f"start_{n.id}")
        starts[k] = new[k] = Node(k, "", "start", cluster=n.cluster)
        e_in.append(Edge(k, n.id))
    for n in sinks:
        k = fresh(f"end_{n.id}")
        ends[k] = new[k] = Node(k, "", "end", cluster=n.cluster)
        e_out.append(Edge(n.id, k))
    d.nodes = {**d.nodes, **starts, **ends}        # after the author's nodes: same cycle breaking
    d.edges = d.edges + e_in + e_out
    return d


# --------------------------------------------------------------------------------------------
# Text metrics
# --------------------------------------------------------------------------------------------

def _char_em(ch: str) -> float:
    if ch in "il.,:;'|!`":
        return 0.28
    if ch in "fjrtI()[]{}- /\\\"":
        return 0.38
    if ch in "mwMW@%":
        return 0.88
    if ord(ch) > 0x2E80:
        return 1.0
    if ch.isupper():
        return 0.66
    if ch.isdigit():
        return 0.56
    return 0.54


def text_width(line: str, size: float) -> float:
    return sum(_char_em(c) for c in line) * size


def wrap(text: str, size: float, max_width: float) -> list[str]:
    lines = []
    for para in text.split("\n"):
        cur = ""
        for word in para.split(" "):
            cand = f"{cur} {word}" if cur else word
            if cur and text_width(cand, size) > max_width:
                lines.append(cur)
                cur = word
            else:
                cur = cand
        lines.append(cur)
    return lines


# --------------------------------------------------------------------------------------------
# Shapes: mermaid shape -> (preset, connection sites side -> (idx, fx, fy), flipH)
# Site indices measured in PowerPoint (ConnectionSiteCount / BeginConnect); fx, fy are the
# site position as a fraction of the box. A callable fraction depends on the box (w, h).
# --------------------------------------------------------------------------------------------

_RECT4 = {"t": (0, .5, 0), "l": (1, 0, .5), "b": (2, .5, 1), "r": (3, 1, .5)}
_OVAL8 = {"t": (0, .5, 0), "l": (2, 0, .5), "b": (4, .5, 1), "r": (6, 1, .5)}
SHAPES = {
    "rect": (MSO_SHAPE.RECTANGLE, _RECT4, False),
    "round": (MSO_SHAPE.ROUNDED_RECTANGLE, _RECT4, False),
    "stadium": (MSO_SHAPE.FLOWCHART_TERMINATOR, _RECT4, False),
    "subroutine": (MSO_SHAPE.FLOWCHART_PREDEFINED_PROCESS, _RECT4, False),
    "diamond": (MSO_SHAPE.FLOWCHART_DECISION, _RECT4, False),
    "hexagon": (MSO_SHAPE.FLOWCHART_PREPARATION, _RECT4, False),
    "circle": (MSO_SHAPE.OVAL, _OVAL8, False),
    "doublecircle": (MSO_SHAPE.OVAL, _OVAL8, False),
    "cylinder": (MSO_SHAPE.FLOWCHART_MAGNETIC_DISK,
                 {"t": (1, .5, 0), "l": (2, 0, .5), "b": (3, .5, 1), "r": (4, 1, .5)}, False),
    "parallelogram": (MSO_SHAPE.FLOWCHART_DATA,
                      {"t": (1, .5, 0), "l": (2, .1, .5), "b": (4, .5, 1), "r": (5, .9, .5)}, False),
    "parallelogram_alt": (MSO_SHAPE.FLOWCHART_DATA,
                          {"t": (1, .5, 0), "l": (5, .1, .5), "b": (4, .5, 1), "r": (2, .9, .5)}, True),
    "trapezoid": (MSO_SHAPE.TRAPEZOID,
                  {"t": (0, .5, 0), "l": (1, lambda w, h: .125 * min(w, h) / w, .5),
                   "b": (2, .5, 1), "r": (3, lambda w, h: 1 - .125 * min(w, h) / w, .5)}, False),
    "trapezoid_alt": (MSO_SHAPE.FLOWCHART_MANUAL_OPERATION,
                      {"t": (0, .5, 0), "l": (1, .1, .5), "b": (2, .5, 1), "r": (3, .9, .5)}, False),
    "flag": (MSO_SHAPE.PENTAGON,
             {"t": (0, lambda w, h: (w - .5 * min(w, h)) / 2 / w, 0), "l": (1, 0, .5),
              "b": (2, lambda w, h: (w - .5 * min(w, h)) / 2 / w, 1), "r": (3, 1, .5)}, False),
    # BPMN render (see to_bpmn)
    "task": (MSO_SHAPE.ROUNDED_RECTANGLE, _RECT4, False),
    "call": (MSO_SHAPE.ROUNDED_RECTANGLE, _RECT4, False),
    "gateway": (MSO_SHAPE.FLOWCHART_DECISION, _RECT4, False),
    "start": (MSO_SHAPE.OVAL, _OVAL8, False),
    "end": (MSO_SHAPE.OVAL, _OVAL8, False),
    "intermediate": (MSO_SHAPE.OVAL, _OVAL8, False),
}
# BPMN gateways and events: a fixed-size symbol (side in em) with its text outside, beside it
OUTSIDE = {"gateway": 2.6, "start": 2.0, "end": 2.0, "intermediate": 2.0}


def site(shape: str, side: str, box: "Box") -> tuple[int, float, float]:
    """Connection site on `side` of a node: (PowerPoint site index, x, y)."""
    idx, fx, fy = SHAPES[shape][1][side]
    if box.core:            # the symbol is a square centred in a box that also holds its text
        box = Box(box.cx - box.core / 2, box.cy - box.core / 2, box.core, box.core)
    fx = fx(box.w, box.h) if callable(fx) else fx
    return idx, box.x + fx * box.w, box.y + fy * box.h


def node_size(shape: str, lines: list[str], fs: float) -> tuple[float, float]:
    """Box (w, h) in points that leaves the text inside the preset's own text area."""
    tw = max((text_width(ln, fs) for ln in lines), default=0.0)
    th = len(lines) * fs * 1.2
    px, py = 0.8 * fs, 0.55 * fs
    w, h = tw + 2 * px, th + 2 * py
    if shape == "diamond":
        w, h = 2 * tw + 2 * px, 2 * th + py
        w = max(w, 1.4 * h)
    elif shape in ("circle", "doublecircle"):
        w = h = max(tw, th) / 0.72 + px
    elif shape == "stadium":
        w += 0.7 * h
    elif shape in ("hexagon", "parallelogram", "parallelogram_alt"):
        w = w / 0.62
    elif shape in ("trapezoid", "trapezoid_alt"):
        w += 0.6 * h
    elif shape == "subroutine":
        w = w / 0.78
    elif shape == "cylinder":
        h = th / 0.5 + py
    elif shape == "flag":
        w += 0.55 * h
    return max(w, 3.5 * fs), max(h, 2.4 * fs)


# --------------------------------------------------------------------------------------------
# Layout
# --------------------------------------------------------------------------------------------

@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float
    core: float = 0.0       # side of the drawn symbol when the text sits outside it (OUTSIDE)

    @property
    def cx(self):
        return self.x + self.w / 2

    @property
    def cy(self):
        return self.y + self.h / 2


@dataclass
class Route:
    edge: Edge
    points: list[tuple[float, float]]   # screen points, orthogonal polyline
    src_site: int | None                # PowerPoint connection-site index, None = not glued
    dst_site: int | None
    label: Box | None = None
    label_lines: list[str] = field(default_factory=list)


@dataclass
class Layout:
    boxes: dict[str, Box]
    lines: dict[str, list[str]]
    routes: list[Route]
    clusters: list[tuple[str, Box, list[str]]]   # outermost first
    width: float
    height: float
    font_size: float
    lanes: list[tuple[str, Box, Box, list[str]]] = field(default_factory=list)   # id, body, header, title


@dataclass
class _LNode:
    key: object          # node id, or ("d", edge index, k) for a dummy
    rs: float            # extent along the rank axis
    cs: float            # extent along the cross axis
    real: bool
    path: tuple
    rank: int
    label: bool = False
    c: float = 0.0


# spacing, in em of the font size
NODESEP, RANKSEP, DUMMYSEP, TRACK, PAD, BACKSEP, LOOP, WRAP_EM = 2.4, 2.9, 0.9, 0.6, 0.9, 1.2, 1.0, 14
OUTSIDE_WRAP_EM, OUTSIDE_GAP, ICON_EM = 9, 0.4, 0.9
OUT = {"TB": "b", "BT": "t", "LR": "r", "RL": "l"}
IN = {"TB": "t", "BT": "b", "LR": "l", "RL": "r"}
SIDE = {"TB": "r", "BT": "r", "LR": "b", "RL": "b"}          # the +cross side
SIDE_NEG = {"TB": "l", "BT": "l", "LR": "t", "RL": "t"}


def _median(vals: list[float]) -> float:
    v = sorted(vals)
    n = len(v)
    return v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2


def _pav(targets: list[float], weights: list[float]) -> list[float]:
    """Weighted isotonic regression (pool adjacent violators): closest non-decreasing fit."""
    blocks: list[list[float]] = []
    for t, w in zip(targets, weights):
        blocks.append([t, w, 1])
        while len(blocks) > 1 and blocks[-2][0] > blocks[-1][0]:
            v2, w2, n2 = blocks.pop()
            v1, w1, n1 = blocks.pop()
            blocks.append([(v1 * w1 + v2 * w2) / (w1 + w2), w1 + w2, n1 + n2])
    out: list[float] = []
    for v, _, n in blocks:
        out += [v] * int(n)
    return out


def _common(a: tuple, b: tuple) -> int:
    n = 0
    while n < min(len(a), len(b)) and a[n] == b[n]:
        n += 1
    return n


def layout(d: Diagram, font_size: float = 12.0) -> Layout:
    fs = font_size
    lfs = 0.85 * fs                                 # edge-label font
    direction = d.direction
    vertical = direction in ("TB", "BT")            # rank axis is screen y
    ids = list(d.nodes)

    # -- sizes ------------------------------------------------------------------------------
    lines = {n: wrap(d.nodes[n].text, fs, WRAP_EM * fs) for n in ids}
    sizes = {n: node_size(d.nodes[n].shape, lines[n], fs) for n in ids}
    cores: dict[str, float] = {}
    for n in ids:           # symbol + its text on the -cross side, mirrored so the symbol is centred
        if (em := OUTSIDE.get(d.nodes[n].shape)) is None:
            continue
        core = cores[n] = em * fs
        lines[n] = wrap(d.nodes[n].text, fs, OUTSIDE_WRAP_EM * fs) if d.nodes[n].text else []
        tw = max((text_width(ln, fs) + 0.3 * fs for ln in lines[n]), default=0.0)
        th = len(lines[n]) * fs * 1.2
        gap = OUTSIDE_GAP * fs if lines[n] else 0.0
        sizes[n] = (core + 2 * (tw + gap), max(core, th)) if vertical else (max(core, tw), core + 2 * (th + gap))
    for n in ids:           # a user icon in the top-left corner: widen both sides, text stays centred
        if d.nodes[n].icon and n not in cores:
            w, h = sizes[n]
            sizes[n] = (w + 2 * (ICON_EM + 0.3) * fs, max(h, (ICON_EM + 0.8) * fs))
    lane_ix = {sg: i for i, sg in enumerate(d.lanes)}
    label_lines = [wrap(e.text, lfs, 0.8 * WRAP_EM * fs) if e.text else [] for e in d.edges]
    label_size = [(max(text_width(ln, lfs) for ln in ll) + 0.5 * lfs, len(ll) * lfs * 1.2 + 0.3 * lfs)
                  if ll else (0.0, 0.0) for ll in label_lines]

    def rc(w, h):
        return (h, w) if vertical else (w, h)

    # -- cycle removal + ranking (longest path, sources pulled down to their children) -------
    edges = d.edges
    loops = {i for i, e in enumerate(edges) if e.src == e.dst}
    out_adj: dict[str, list[int]] = {n: [] for n in ids}
    indeg = {n: 0 for n in ids}
    for i, e in enumerate(edges):
        if i not in loops:
            out_adj[e.src].append(i)
            indeg[e.dst] += 1
    reversed_: set[int] = set()
    state: dict[str, int] = {}
    for root in ids:                    # declaration order: the edge written last closes a cycle
        if root in state:
            continue
        state[root] = 1
        stack = [(root, iter(out_adj[root]))]
        while stack:
            n, it = stack[-1]
            for i in it:
                v = edges[i].dst
                if state.get(v) == 1:
                    reversed_.add(i)
                elif v not in state:
                    state[v] = 1
                    stack.append((v, iter(out_adj[v])))
                    break
            else:
                state[n] = 2
                stack.pop()

    def minlen(i):
        e = edges[i]
        return max(e.minlen, 2) if e.text and i not in reversed_ else e.minlen

    succ: dict[str, list[tuple[str, int]]] = {n: [] for n in ids}
    pred: dict[str, list[tuple[str, int]]] = {n: [] for n in ids}
    for i, e in enumerate(edges):
        if i in loops:
            continue
        u, v = (e.dst, e.src) if i in reversed_ else (e.src, e.dst)
        succ[u].append((v, minlen(i)))
        pred[v].append((u, minlen(i)))
    remaining = {n: len(pred[n]) for n in ids}
    topo = [n for n in ids if not remaining[n]]
    for n in topo:
        for v, _ in succ[n]:
            remaining[v] -= 1
            if not remaining[v]:
                topo.append(v)
    rank = {n: 0 for n in ids}
    for n in topo:
        for v, ml in succ[n]:
            rank[v] = max(rank[v], rank[n] + ml)
    for n in reversed(topo):
        if not pred[n] and succ[n]:
            rank[n] = min(rank[v] - ml for v, ml in succ[n])
    low = min(rank.values(), default=0)
    rank = {n: r - low for n, r in rank.items()}

    through: set[int] = set()        # back edges threaded through the layers
    while True:
        # -- layout graph: real nodes + dummy chains for forward edges ---------------------------
        paths = {n: d.path(n) for n in ids}
        ln: dict[object, _LNode] = {}
        for n in ids:
            rs, cs = rc(*sizes[n])
            ln[n] = _LNode(n, rs, cs, True, paths[n], rank[n])
        chains: dict[int, list] = {}         # edge -> [upper node, dummies..., lower node]
        for i, e in enumerate(edges):
            if i in loops or (i in reversed_ and i not in through):
                continue
            u, v = (e.dst, e.src) if i in reversed_ else (e.src, e.dst)
            span = rank[v] - rank[u]
            common = paths[u][:_common(paths[u], paths[v])]
            if lane_ix and not common:      # a flow between lanes runs in the lower node's lane
                common = paths[v][:1]
            keys = [u]
            for k in range(1, span):
                is_label = bool(e.text) and k == span // 2
                rs, cs = rc(*label_size[i]) if is_label else (0.0, 0.0)
                key = ("d", i, k)
                ln[key] = _LNode(key, rs, cs, False, common, rank[u] + k, is_label)
                keys.append(key)
            keys.append(v)
            chains[i] = keys
        down: dict[object, list] = {k: [] for k in ln}
        up: dict[object, list] = {k: [] for k in ln}
        for keys in chains.values():
            for a, b in zip(keys, keys[1:]):
                down[a].append(b)
                up[b].append(a)

        n_ranks = max((x.rank for x in ln.values()), default=0) + 1
        layers: list[list] = [[] for _ in range(n_ranks)]
        seen: set = set()
        for root in [n for n in ids if not up[n]] + ids:     # DFS gives the initial order
            stack = [root]
            while stack:
                k = stack.pop()
                if k in seen:
                    continue
                seen.add(k)
                layers[ln[k].rank].append(k)
                stack.extend(reversed(down[k]))

        def lane_of(k) -> int:
            return lane_ix[ln[k].path[0]] if lane_ix else 0

        layers = [sorted(layer, key=lane_of) for layer in layers]      # lanes in declaration order

        # -- ordering: barycenter sweeps, subgraph members kept contiguous -----------------------
        pos = {k: i for layer in layers for i, k in enumerate(layer)}

        def crossings() -> int:
            total = 0
            for layer in layers[:-1]:
                segs = [(pos[a], pos[b]) for a in layer for b in down[a]]
                total += sum(1 for x in range(len(segs)) for y in range(x + 1, len(segs))
                             if (segs[x][0] - segs[y][0]) * (segs[x][1] - segs[y][1]) < 0)
            return total

        def sort_layer(layer: list, neigh) -> list:
            bary = {k: (sum(pos[n] for n in neigh[k]) / len(neigh[k]) if neigh[k] else float(pos[k]))
                    for k in layer}
            depth = max((len(ln[k].path) for k in layer), default=0)
            means = {}
            for k in layer:
                p = ln[k].path
                for lvl in range(len(p)):
                    means.setdefault(p[:lvl + 1], []).append(bary[k])

            def key(k):
                p = ln[k].path
                out = [(lane_of(k), "")]
                for lvl in range(depth):
                    if len(p) > lvl:
                        g = p[:lvl + 1]
                        out.append((sum(means[g]) / len(means[g]), "/".join(g)))
                    else:
                        out.append((bary[k], ""))
                out.append((bary[k], str(pos[k])))
                return out

            return sorted(layer, key=key)

        best, best_layers = crossings(), [list(la) for la in layers]
        for it in range(24):
            rng = range(1, n_ranks) if it % 2 == 0 else range(n_ranks - 2, -1, -1)
            for r in rng:
                layers[r] = sort_layer(layers[r], up if it % 2 == 0 else down)
                pos.update({k: i for i, k in enumerate(layers[r])})
            cur = crossings()
            if cur < best:
                best, best_layers = cur, [list(la) for la in layers]
        layers = best_layers
        pos = {k: i for layer in layers for i, k in enumerate(layer)}

        # -- cross coordinate: median alignment + isotonic packing -------------------------------
        title_h = fs * 1.5
        pad = PAD * fs

        def sep(a, b) -> float:
            na, nb = ln[a], ln[b]
            solid = (na.real or na.label) and (nb.real or nb.label)
            s = (na.cs + nb.cs) / 2 + (NODESEP if solid else DUMMYSEP) * fs
            c = _common(na.path, nb.path)
            s += ((len(na.path) - c) + (len(nb.path) - c)) * pad
            if not vertical:                               # cluster titles sit above, on the c axis
                s += (len(nb.path) - c) * title_h
            return s

        offsets = []
        for layer in layers:
            off = [0.0]
            for a, b in zip(layer, layer[1:]):
                off.append(off[-1] + sep(a, b))
            offsets.append(off)
            for k, o in zip(layer, off):
                ln[k].c = o

        # the out/in site of a node may be off-centre on the cross axis (flag shape)
        def c_off(n, side):
            w, h = sizes[n]
            _, x, y = site(d.nodes[n].shape, side, Box(-w / 2, -h / 2, w, h, cores.get(n, 0.0)))
            return x if vertical else y

        def anchor(k, other):
            """Cross coordinate an edge leaves/enters `k` from, seen from neighbour `other`."""
            node = ln[k]
            if not node.real:
                return node.c
            return node.c + c_off(k, OUT[direction] if ln[other].rank > node.rank else IN[direction])

        def place(r, neigh):
            layer = layers[r]
            targets, weights = [], []
            for k in layer:
                ns = neigh[k]
                if ns:
                    want = _median([anchor(n, k) for n in ns])
                    if ln[k].real:
                        want -= anchor(k, ns[0]) - ln[k].c
                    targets.append(want)
                else:
                    targets.append(ln[k].c)
                weights.append(1.0 if ln[k].real else 4.0)
            ys = _pav([t - o for t, o in zip(targets, offsets[r])], weights)
            for k, y, o in zip(layer, ys, offsets[r]):
                ln[k].c = y + o

        for it in range(10):
            if it % 2 == 0:
                for r in range(1, n_ranks):
                    place(r, up)
            else:
                for r in range(n_ranks - 2, -1, -1):
                    place(r, down)

        def try_move(k, c) -> bool:
            """Move `k` to cross coordinate `c` if its layer neighbours leave room."""
            layer, p = layers[ln[k].rank], pos[k]
            if p > 0 and c - ln[layer[p - 1]].c < sep(layer[p - 1], k) - 0.01:
                return False
            if p + 1 < len(layer) and ln[layer[p + 1]].c - c < sep(k, layer[p + 1]) - 0.01:
                return False
            ln[k].c = c
            return True

        # straighten dummy chains: one lane, ideally in line with the source or the target
        for keys in chains.values():
            dummies = keys[1:-1]
            if not dummies:
                continue
            for lane in (anchor(keys[-1], keys[-2]), anchor(keys[0], keys[1]), _median([ln[k].c for k in dummies])):
                old = [ln[k].c for k in dummies]
                if all(try_move(k, lane) for k in dummies):
                    break
                for k, c in zip(dummies, old):
                    ln[k].c = c

        # 1-to-1 links between real nodes: put the pair in line when there is room
        for rng, nb, back in ((range(1, n_ranks), up, down), (range(n_ranks - 2, -1, -1), down, up)):
            for r in rng:
                for k in layers[r]:
                    if ln[k].real and len(nb[k]) == 1 and len(back[nb[k][0]]) == 1 and ln[nb[k][0]].real:
                        o = nb[k][0]
                        try_move(k, ln[k].c + anchor(o, k) - anchor(k, o))

        # back edges go round the outside of the layers, on the side where both of their nodes
        # are outermost; when there is no such side, lay out again with them threaded through
        # the layers like forward edges (dummy slots reserved), drawn upwards
        # (with lanes: outermost in its own lane, and only for an edge that stays in one lane)
        def outermost(n, sign):
            return all(not ln[k].real or lane_of(k) != lane_of(n) or (ln[k].c - ln[n].c) * sign <= 0
                       for k in layers[rank[n]])

        back_side: dict[int, int] = {}
        blocked = set()
        for i in reversed_ - through:
            e = edges[i]
            sign = None if lane_of(e.src) != lane_of(e.dst) else \
                next((sg for sg in (1, -1) if outermost(e.src, sg) and outermost(e.dst, sg)), None)
            if sign is None:
                blocked.add(i)
            else:
                back_side[i] = sign
        if blocked:
            through |= blocked
            continue

        # cross coordinate of each back edge's side run; overlapping ones are stacked outwards
        side_c: dict[int, float] = {}
        side_runs: list[tuple[int, int, int, int, float]] = []
        for i in sorted(back_side, key=lambda i: abs(rank[edges[i].src] - rank[edges[i].dst])):
            e, sign = edges[i], back_side[i]
            lo, hi = sorted((rank[e.src], rank[e.dst]))

            def reach(k):       # how far a node and the subgraph boxes around it extend on that side
                depth = len(ln[k].path)
                titles = (depth - bool(lane_ix)) * title_h if not vertical and sign < 0 else 0.0
                return sign * ln[k].c + ln[k].cs / 2 + depth * pad + titles

            run = sign * (max(reach(k) for k in ln if lo <= ln[k].rank <= hi and lane_of(k) == lane_of(e.src))
                          + BACKSEP * fs)
            for a, b, sg, ls, lc in side_runs:
                if sg == sign and ls == lane_of(e.src) and a <= hi and lo <= b:
                    run = sign * max(sign * run, sign * lc + BACKSEP * fs)
            lcs = rc(*label_size[i])[1] if e.text else 0.0
            side_c[i] = run + sign * lcs / 2
            side_runs.append((lo, hi, sign, lane_of(e.src), run + sign * lcs))

        # lanes: shift each lane as a whole so their bands sit side by side, in order
        bands: list[tuple[str, float, float]] = []
        if lane_ix:
            ext: dict[int, list[float]] = {}
            for k, node in ln.items():
                m = len(node.path) * pad
                lo_hi = ext.setdefault(lane_of(k), [node.c - node.cs / 2 - m, node.c + node.cs / 2 + m])
                lo_hi[0] = min(lo_hi[0], node.c - node.cs / 2 - m)
                lo_hi[1] = max(lo_hi[1], node.c + node.cs / 2 + m)
            for i, c in side_c.items():
                sign, (lw, lh) = back_side[i], label_size[i]
                reach_ = c + sign * (rc(lw, lh)[1] / 2 + pad)
                lo_hi = ext[lane_of(edges[i].src)]
                lo_hi[0], lo_hi[1] = min(lo_hi[0], reach_), max(lo_hi[1], reach_)
            for i in loops:
                n = edges[i].src
                extra = LOOP * fs + (rc(*label_size[i])[1] + 0.2 * fs if edges[i].text else 0.0) + pad
                ext[lane_of(n)][1] = max(ext[lane_of(n)][1], ln[n].c + ln[n].cs / 2 + extra)
            delta, end = {}, 0.0
            for sg, ix in lane_ix.items():
                lo, hi = ext.get(ix, [end, end])
                title = d.subgraphs[sg]
                want = min(text_width(title.title, fs) + (ICON_EM + 0.4) * fs * bool(title.icon),
                           WRAP_EM * fs) + 2 * pad
                if hi - lo < want:
                    lo, hi = (lo + hi - want) / 2, (lo + hi + want) / 2
                delta[ix] = end - lo
                bands.append((sg, end, end + hi - lo))
                end += hi - lo
            for k in ln:
                ln[k].c += delta[lane_of(k)]
            for i in side_c:
                side_c[i] += delta[lane_of(edges[i].src)]

        # -- tracks: horizontal jogs of different edges in one gap never share a line ------------
        snap = 0.15 * fs
        out_deg = {n: 0 for n in ids}
        in_deg = {n: 0 for n in ids}
        for keys in chains.values():
            out_deg[keys[0]] += 1
            in_deg[keys[-1]] += 1
        jogs: dict[int, list] = {}            # edge -> [(gap, from_c, to_c, key)]
        intervals: dict[int, dict] = {}       # gap -> key -> [lo, hi]
        for i, keys in chains.items():
            cur = anchor(keys[0], keys[1])
            jogs[i] = []
            for j in range(1, len(keys)):
                a, b = keys[j - 1], keys[j]
                to = anchor(b, a)
                if abs(to - cur) <= snap:
                    continue
                if i in reversed_:          # threaded back edge: its own track, never shared
                    key = ("e", i, j)
                elif ln[a].real and ln[b].real:
                    key = ("s", a) if out_deg[a] >= in_deg[b] else ("t", b)
                elif ln[a].real:
                    key = ("s", a)
                elif ln[b].real:
                    key = ("t", b)
                else:
                    key = ("e", i, j)
                g = ln[a].rank
                jogs[i].append((g, cur, to, key))
                iv = intervals.setdefault(g, {}).setdefault(key, [min(cur, to), max(cur, to)])
                iv[0], iv[1] = min(iv[0], cur, to), max(iv[1], cur, to)
                cur = to
        track_of: dict[tuple, int] = {}
        n_tracks = [0] * n_ranks
        for g, ivs in intervals.items():
            ends: list[float] = []
            for key, (lo, hi) in sorted(ivs.items(), key=lambda kv: kv[1][0]):
                t = next((t for t, e in enumerate(ends) if e < lo - 0.5 * fs), None)
                if t is None:
                    t = len(ends)
                    ends.append(hi)
                ends[t] = hi
                track_of[(g, key)] = t
            n_tracks[g] = len(ends)

        # -- rank coordinate ---------------------------------------------------------------------
        sg_ranks: dict[str, list[int]] = {}
        for n in ids:
            for sg in paths[n]:
                if sg not in lane_ix:       # a lane's title is in its header, not between layers
                    sg_ranks.setdefault(sg, []).append(rank[n])
        layer_size = [max((ln[k].rs for k in layer), default=0.0) for layer in layers]
        gaps = []
        for g in range(n_ranks - 1):
            starts = sum(1 for rs in sg_ranks.values() if min(rs) == g + 1)
            ends_ = sum(1 for rs in sg_ranks.values() if max(rs) == g)
            extra = (starts + ends_) * pad
            if vertical:
                extra += title_h * (starts if direction == "TB" else ends_)
            gaps.append(max(RANKSEP * fs, (n_tracks[g] + 1) * TRACK * fs) + extra)
        r_center = [layer_size[0] / 2] if n_ranks else []
        for g in range(n_ranks - 1):
            r_center.append(r_center[-1] + layer_size[g] / 2 + gaps[g] + layer_size[g + 1] / 2)
        r_max = (r_center[-1] + layer_size[-1] / 2) if n_ranks else 0.0

        def track_r(g, key):
            start = r_center[g] + layer_size[g] / 2
            return start + (track_of[(g, key)] + 1) * gaps[g] / (n_tracks[g] + 1)

        def to_screen(r, c):
            return {"TB": (c, r), "BT": (c, r_max - r), "LR": (r, c), "RL": (r_max - r, c)}[direction]

        def to_abs(x, y):
            return {"TB": (y, x), "BT": (r_max - y, x), "LR": (x, y), "RL": (r_max - x, y)}[direction]

        boxes = {}
        for n in ids:
            w, h = sizes[n]
            x, y = to_screen(r_center[rank[n]], ln[n].c)
            boxes[n] = Box(x - w / 2, y - h / 2, w, h, cores.get(n, 0.0))

        def site_abs(n, side):
            idx, x, y = site(d.nodes[n].shape, side, boxes[n])
            return idx, to_abs(x, y)

        # -- routes ------------------------------------------------------------------------------
        routes: list[Route] = []
        for i, keys in chains.items():
            e = edges[i]
            si, (sr, sc) = site_abs(keys[0], OUT[direction])
            ti, (er, ec) = site_abs(keys[-1], IN[direction])
            pts = [(sr, sc)]
            for g, frm, to, key in jogs[i]:
                tr = track_r(g, key)
                pts += [(tr, frm), (tr, to)]
            pts.append((er, pts[-1][1] if abs(ec - pts[-1][1]) <= snap else ec))
            screen = [to_screen(r, c) for r, c in pts]
            route = Route(e, screen[::-1], ti, si) if i in reversed_ else Route(e, screen, si, ti)
            lab = next((k for k in keys[1:-1] if ln[k].label), None)
            lw, lh = label_size[i]
            if lab is not None:
                x, y = to_screen(r_center[ln[lab].rank], ln[lab].c)
            elif e.text:        # a threaded back edge too short for a label slot: longest segment
                a, b = max(zip(screen, screen[1:]), key=lambda s: abs(s[0][0] - s[1][0]) + abs(s[0][1] - s[1][1]))
                x, y = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
            if e.text:
                route.label, route.label_lines = Box(x - lw / 2, y - lh / 2, lw, lh), label_lines[i]
            routes.append(route)

        # back edges: out of the side of the lower node, up a lane beside everything, into the
        # side of the upper one; lanes of overlapping back edges are stacked outwards
        for i, run in side_c.items():
            e, sign = edges[i], back_side[i]
            side = SIDE[direction] if sign > 0 else SIDE_NEG[direction]
            si, (sr, sc) = site_abs(e.src, side)
            ti, (er, ec) = site_abs(e.dst, side)
            lw, lh = label_size[i]
            route = Route(e, [to_screen(*p) for p in [(sr, sc), (sr, run), (er, run), (er, ec)]], si, ti)
            if e.text:
                x, y = to_screen((sr + er) / 2, run)
                route.label, route.label_lines = Box(x - lw / 2, y - lh / 2, lw, lh), label_lines[i]
            routes.append(route)
        break

    for i in sorted(loops):
        e = edges[i]
        si, (sr, sc) = site_abs(e.src, SIDE[direction])
        ti, (er, ec) = site_abs(e.src, IN[direction])
        m = LOOP * fs
        top = er - m
        pts = [(sr, sc), (sr, sc + m), (top, sc + m), (top, ec), (er, ec)]
        route = Route(e, [to_screen(r, c) for r, c in pts], si, ti)
        if e.text:
            lw, lh = label_size[i]
            x, y = to_screen((sr + top) / 2, sc + m + rc(lw, lh)[1] / 2 + 0.2 * fs)
            route.label, route.label_lines = Box(x - lw / 2, y - lh / 2, lw, lh), label_lines[i]
        routes.append(route)
    routes.sort(key=lambda r: edges.index(r.edge))

    # -- subgraph boxes, innermost first so parents can wrap them -----------------------------
    cl_boxes: dict[str, Box] = {}
    order = sorted((s for s in d.subgraphs if s not in lane_ix), key=lambda s: -len(_ancestors(d, s)))
    for sid in order:
        parts = [boxes[n] for n in ids if d.nodes[n].cluster == sid]
        parts += [cl_boxes[s] for s in d.subgraphs if d.subgraphs[s].parent == sid and s in cl_boxes]
        parts += [r.label for r in routes if r.label and sid in paths[r.edge.src]
                  and sid in paths[r.edge.dst] and r.edge.src != r.edge.dst]
        if not parts:
            continue
        x0 = min(b.x for b in parts) - pad
        y0 = min(b.y for b in parts) - pad - title_h
        x1 = max(b.x + b.w for b in parts) + pad
        y1 = max(b.y + b.h for b in parts) + pad
        cl_boxes[sid] = Box(x0, y0, x1 - x0, y1 - y0)
    drawn = sorted(cl_boxes, key=lambda s: (len(_ancestors(d, s)), -cl_boxes[s].w * cl_boxes[s].h))
    clusters = [(s, cl_boxes[s], [d.subgraphs[s].title]) for s in drawn]

    # -- lanes: one band each across the whole diagram, a header strip at the start of the flow --
    lane_boxes: list[tuple[str, Box, Box, list[str]]] = []
    if bands:
        allb = list(boxes.values()) + list(cl_boxes.values()) + [r.label for r in routes if r.label]
        pts = [p for r in routes for p in r.points]
        # extent along the rank axis (screen y when vertical, x otherwise)
        r0 = min([b.y if vertical else b.x for b in allb] + [p[1] if vertical else p[0] for p in pts]) - pad
        r1 = max([b.y + b.h if vertical else b.x + b.w for b in allb]
                 + [p[1] if vertical else p[0] for p in pts]) + pad
        titles = {}
        for sg, lo, hi in bands:
            icon = (ICON_EM + 0.4) * fs if d.subgraphs[sg].icon else 0.0
            titles[sg] = wrap(d.subgraphs[sg].title, fs, max(hi - lo - pad - icon, 3 * fs))
        head = max(len(t) for t in titles.values()) * fs * 1.2 + 0.8 * fs
        for sg, lo, hi in bands:
            if vertical:
                body, header = Box(lo, r0, hi - lo, r1 - r0), Box(lo, r0 - head, hi - lo, head)
            else:
                body, header = Box(r0, lo, r1 - r0, hi - lo), Box(r0 - head, lo, head, hi - lo)
            lane_boxes.append((sg, body, header, titles[sg]))

    # -- normalise to the origin -------------------------------------------------------------
    xs, ys = [], []
    for b in list(boxes.values()) + list(cl_boxes.values()) + [r.label for r in routes if r.label] + \
            [b for _, body, header, _ in lane_boxes for b in (body, header)]:
        xs += [b.x, b.x + b.w]
        ys += [b.y, b.y + b.h]
    for r in routes:
        xs += [p[0] for p in r.points]
        ys += [p[1] for p in r.points]
    x0, y0 = (min(xs), min(ys)) if xs else (0.0, 0.0)
    width, height = ((max(xs) - x0), (max(ys) - y0)) if xs else (0.0, 0.0)

    def shift(b):
        return replace(b, x=b.x - x0, y=b.y - y0) if b else None

    for r in routes:
        r.points = [(x - x0, y - y0) for x, y in r.points]
        r.label = shift(r.label)
    return Layout({n: shift(b) for n, b in boxes.items()}, lines, routes,
                  [(s, shift(b), t) for s, b, t in clusters], width, height, fs,
                  [(s, shift(b), shift(h), t) for s, b, h, t in lane_boxes])


# --------------------------------------------------------------------------------------------
# Connector geometry
# --------------------------------------------------------------------------------------------

_ROT = {(1, 0): 0, (0, 1): 5400000, (-1, 0): 10800000, (0, -1): 16200000}
# a straight edge is a flat elbow (bentConnector3 of height 0): it stays orthogonal when
# PowerPoint re-routes it after a node is moved, where a straightConnector1 turns diagonal
_PRESET = {1: "bentConnector3", 2: "bentConnector2", 3: "bentConnector3",
           4: "bentConnector4", 5: "bentConnector5"}


def simplify(points: list[tuple[float, float]], eps: float = 0.01) -> list[tuple[float, float]]:
    """Drop zero-length segments and merge collinear ones."""
    out: list[tuple[float, float]] = []
    for p in points:
        if out and abs(p[0] - out[-1][0]) < eps and abs(p[1] - out[-1][1]) < eps:
            continue
        if len(out) >= 2:
            a, b = out[-2], out[-1]
            if (abs(a[0] - b[0]) < eps and abs(b[0] - p[0]) < eps) or \
               (abs(a[1] - b[1]) < eps and abs(b[1] - p[1]) < eps):
                out[-1] = p
                continue
        out.append(p)
    return out


def connector_geometry(points: list[tuple[float, float]], min_extent: float = 0.1) -> dict:
    """Preset, rotation, flip, box and adjust values of a connector drawing `points`.

    `points` is an orthogonal polyline (1 to 5 segments). A bentConnectorN draws, in its own
    box, (0,0) -> (x1,0) -> (x1,y2) -> (x3,y2) ... -> (w,h), the first segment along +x;
    `rot` (a multiple of 90 degrees) turns +x onto the first segment's direction and `flipV`
    puts the end on the right side, so the drawn path is exactly `points`.
    """
    n = len(points) - 1
    if not 1 <= n <= 5:
        raise ValueError(f"a connector draws 1 to 5 segments, not {n}")
    (sx, sy), (ex, ey) = points[0], points[-1]
    dx, dy = points[1][0] - sx, points[1][1] - sy
    first = (1 if dx > 0 else -1, 0) if abs(dx) >= abs(dy) else (0, 1 if dy > 0 else -1)
    a = (ex - sx) * first[0] + (ey - sy) * first[1]
    ux = first if a >= 0 or n == 1 else (-first[0], -first[1])
    uy_rot = (-ux[1], ux[0])                         # where the rotation sends local +y
    b = (ex - sx) * uy_rot[0] + (ey - sy) * uy_rot[1]
    flip = b < 0
    uy = (-uy_rot[0], -uy_rot[1]) if flip else uy_rot
    w, h = abs(a), abs(b)
    if n > 1:
        w, h = max(w, min_extent), max(h, min_extent)
    local = [((x - sx) * ux[0] + (y - sy) * ux[1], (x - sx) * uy[0] + (y - sy) * uy[1])
             for x, y in points]
    adj = [round(100000 * (local[k][0] / w if k % 2 else local[k][1] / h)) for k in range(1, n - 1)] \
        if n > 1 else [50000]
    cx = sx + ux[0] * w / 2 + uy[0] * h / 2
    cy = sy + ux[1] * w / 2 + uy[1] * h / 2
    return dict(prst=_PRESET[n], rot=_ROT[ux], flipV=flip, x=cx - w / 2, y=cy - h / 2, w=w, h=h, adj=adj)


# --------------------------------------------------------------------------------------------
# Rendering
# --------------------------------------------------------------------------------------------

# node_*: boxes; accent: event outlines; text: labels outside the boxes
PALETTES = {
    "purple": {
        "node_fill": "5236AB", "node_line": "3A2680", "node_text": "FFFFFF", "accent": "5236AB",
        "text": "200A58", "edge": "200A58", "cluster_fill": "F2F1F9", "cluster_line": "A8A0D6",
        "cluster_text": "200A58", "label_fill": "FFFFFF",
    },
    "slate": {
        "node_fill": "F1F5F9", "node_line": "64748B", "node_text": "0F172A", "accent": "64748B",
        "text": "0F172A", "edge": "475569", "cluster_fill": "F8FAFC", "cluster_line": "94A3B8",
        "cluster_text": "334155", "label_fill": "FFFFFF",
    },
    "theme": {  # recolours with the deck it is pasted into
        "node_fill": MSO_THEME_COLOR.BACKGROUND_1, "node_line": MSO_THEME_COLOR.ACCENT_1,
        "node_text": MSO_THEME_COLOR.TEXT_1, "accent": MSO_THEME_COLOR.ACCENT_1,
        "text": MSO_THEME_COLOR.TEXT_1, "edge": MSO_THEME_COLOR.TEXT_1,
        "cluster_fill": MSO_THEME_COLOR.BACKGROUND_2, "cluster_line": MSO_THEME_COLOR.TEXT_2,
        "cluster_text": MSO_THEME_COLOR.TEXT_2, "label_fill": MSO_THEME_COLOR.BACKGROUND_1,
    },
}
AUTHOR_ENV = "MERMAID2PPTX_AUTHOR"      # default document author, when --author is not given


def _contrast(rgb: str) -> str:
    r, g, b = (int(rgb[k:k + 2], 16) for k in (0, 2, 4))
    return "0F172A" if 0.299 * r + 0.587 * g + 0.114 * b > 150 else "FFFFFF"


def palette(color: str) -> dict:
    """A named palette (PALETTES), or the slate one with boxes in the colour `#RRGGBB`."""
    if color.lower() in PALETTES:
        return PALETTES[color.lower()]
    rgb = parse_color(color)
    if rgb is None:
        raise ValueError(f"unknown colour {color!r}: use {', '.join(PALETTES)} or #RRGGBB")
    dark = "".join(f"{int(int(rgb[k:k + 2], 16) * 0.7):02X}" for k in (0, 2, 4))
    return {**PALETTES["slate"], "node_fill": rgb, "node_line": dark, "node_text": _contrast(rgb), "accent": rgb}
_MARKER_XML = {"arrow": "triangle", "circle": "oval", "cross": "diamond"}
SLIDE_SIZES = {"16:9": (960.0, 540.0), "4:3": (720.0, 540.0)}
MARGIN = 24.0


@dataclass
class Options:
    font_size: float = 12.0
    font: str | None = None
    color: str = "purple"               # a PALETTES name or #RRGGBB
    render: str = "bpmn"                # bpmn | mermaid
    events: bool = True                 # bpmn: add missing start/end events
    lanes: bool = True                  # bpmn: top-level subgraphs as swim lanes
    author: str | None = None           # None: $MERMAID2PPTX_AUTHOR, else empty
    group: bool = True
    fit: bool = True
    aspect: str = "16:9"
    direction: str | None = None


def _set_color(color_format, spec):
    if isinstance(spec, str):
        color_format.rgb = RGBColor.from_string(spec)
    else:
        color_format.theme_color = spec


def _strip_style(shape):
    """Drop the theme style reference: no inherited shadow, fill or font colour."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _dashed(style: dict[str, str]) -> bool:
    """stroke-dasharray with a non-zero dash (`0` and `none` mean a solid line)."""
    return bool(re.search(r"[1-9]", style.get("stroke-dasharray", "")))


def _px_to_pt(v: str) -> float | None:
    m = re.match(r"([\d.]+)\s*(px|pt)?", v)
    if not m:
        return None
    return float(m[1]) * (0.75 if (m[2] or "px") == "px" else 1.0)


def _write_text(tf, lines, size, color, font, bold=False, align=PP_ALIGN.CENTER):
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = "\v".join(lines)
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        _set_color(run.font.color, color)
        if font:
            run.font.name = font


def render(d: Diagram, lay: Layout, opts: Options) -> Presentation:
    pal = palette(opts.color)
    prs = Presentation()
    now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    cp = prs.core_properties         # replace the built-in deck's own metadata
    cp.author = cp.last_modified_by = opts.author if opts.author is not None else os.environ.get(AUTHOR_ENV, "")
    cp.title, cp.subject, cp.comments, cp.keywords, cp.category = "", "", "", "", ""
    cp.created = cp.modified = now
    cp.revision = 1
    sw, sh = SLIDE_SIZES[opts.aspect]
    if opts.fit:
        s = min(1.0, (sw - 2 * MARGIN) / max(lay.width, 1), (sh - 2 * MARGIN) / max(lay.height, 1))
    else:
        s = 1.0
        sw, sh = max(lay.width + 2 * MARGIN, 72), max(lay.height + 2 * MARGIN, 72)
    prs.slide_width, prs.slide_height = Emu(round(sw * EMU_PER_PT)), Emu(round(sh * EMU_PER_PT))
    ox, oy = (sw - lay.width * s) / 2, (sh - lay.height * s) / 2
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    grp = slide.shapes.add_group_shape() if opts.group else None
    shapes = grp.shapes if grp is not None else slide.shapes
    if grp is not None:
        grp.name = "Mermaid diagram"

    def E(v):
        return int(round(v * s * EMU_PER_PT))

    def X(v):
        return int(round((ox + v * s) * EMU_PER_PT))

    def Y(v):
        return int(round((oy + v * s) * EMU_PER_PT))

    fs = max(5.0, round(lay.font_size * s * 2) / 2)
    lfs = max(5.0, round(0.85 * lay.font_size * s * 2) / 2)
    font = opts.font

    def user_icon(x, y, size, color, name):
        """BPMN user marker: a head over rounded shoulders, top-left corner at (x, y)."""
        for preset, (fx, fy, fw, fh) in ((MSO_SHAPE.OVAL, (.28, 0, .44, .44)),
                                         (MSO_SHAPE.ROUND_2_SAME_RECTANGLE, (.06, .5, .88, .5))):
            ic = shapes.add_shape(preset, X(x + fx * size), Y(y + fy * size), E(fw * size), E(fh * size))
            _strip_style(ic)
            ic.name = name
            if preset == MSO_SHAPE.ROUND_2_SAME_RECTANGLE:
                ic.adjustments[0] = 0.5
            ic.fill.solid()
            _set_color(ic.fill.fore_color, color)
            ic.line.fill.background()

    lf = lay.font_size
    icon = ICON_EM * lf
    vertical = d.direction in ("TB", "BT")
    for sid, body, header, title in lay.lanes:
        sg = d.subgraphs[sid]
        sp = shapes.add_shape(MSO_SHAPE.RECTANGLE, X(body.x), Y(body.y), E(body.w), E(body.h))
        _strip_style(sp)
        sp.name = f"lane {sid}"
        _set_color_fill(sp, sg.style.get("fill"), pal["cluster_fill"])
        sp.line.width = Pt(0.75)
        _set_color(sp.line.color, parse_color(sg.style.get("stroke", "")) or pal["cluster_line"])
        hd = shapes.add_shape(MSO_SHAPE.RECTANGLE, X(header.x), Y(header.y), E(header.w), E(header.h))
        _strip_style(hd)
        hd.name = f"lane header {sid}"
        _set_color_fill(hd, None, pal["node_fill"])
        hd.line.width = Pt(0.75)
        _set_color(hd.line.color, pal["node_line"])
        tf = hd.text_frame
        tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.MIDDLE
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = Pt(0.2 * fs)
        if not vertical:
            tf._txBody.find(qn("a:bodyPr")).set("vert", "vert270")
        _write_text(tf, title, fs, pal["node_text"], font, bold=True)
        if sg.icon:     # beside the title's first line, on its reading side
            tw = max(text_width(ln, lf) for ln in title)
            th = len(title) * lf * 1.2
            if vertical:
                user_icon(header.cx - tw / 2 - icon - 0.3 * lf, header.cy - th / 2 + 0.1 * lf, icon,
                          pal["node_text"], f"icon {sid}")
            else:
                user_icon(header.cx - th / 2 + 0.1 * lf, header.cy + tw / 2 + 0.3 * lf, icon,
                          pal["node_text"], f"icon {sid}")

    for sid, b, title in lay.clusters:
        sp = shapes.add_shape(MSO_SHAPE.RECTANGLE, X(b.x), Y(b.y), E(b.w), E(b.h))
        _strip_style(sp)
        sp.name = f"subgraph {sid}"
        style = d.subgraphs[sid].style
        _set_color_fill(sp, style.get("fill"), pal["cluster_fill"])
        sp.line.width = Pt(_px_to_pt(style.get("stroke-width", "")) or 0.75)
        _set_color(sp.line.color, parse_color(style.get("stroke", "")) or pal["cluster_line"])
        if opts.render == "bpmn":       # a BPMN group: dash-dot border
            sp.line.dash_style = MSO_LINE_DASH_STYLE.LONG_DASH_DOT
        tf = sp.text_frame
        tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.TOP
        tf.margin_top = tf.margin_bottom = Pt(0.25 * fs)
        tf.margin_left = tf.margin_right = Pt(0.4 * fs)
        _write_text(tf, title, fs, pal["cluster_text"], font, bold=True,   # BPMN: label top-left
                    align=PP_ALIGN.LEFT if opts.render == "bpmn" else PP_ALIGN.CENTER)

    ids: dict[str, int] = {}
    for n, b in lay.boxes.items():
        node = d.nodes[n]
        preset, _, flip_h = SHAPES[node.shape]
        shape_box = Box(b.cx - b.core / 2, b.cy - b.core / 2, b.core, b.core) if b.core else b
        sp = shapes.add_shape(preset, X(shape_box.x), Y(shape_box.y), E(shape_box.w), E(shape_box.h))
        _strip_style(sp)
        sp.name = f"node {n}"
        ids[n] = sp.shape_id
        if flip_h:
            sp._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
        if node.shape in ("task", "call"):
            sp.adjustments[0] = 0.12
        style = d.node_style(node)
        event = node.shape in ("start", "end", "intermediate")
        fill = parse_color(style.get("fill", ""))
        _set_color_fill(sp, style.get("fill"), pal["label_fill"] if event else pal["node_fill"])
        stroke = parse_color(style.get("stroke", ""))
        _set_color(sp.line.color, stroke or (pal["accent"] if event else pal["node_line"]))
        width = _px_to_pt(style.get("stroke-width", "")) or \
            {"doublecircle": 3, "call": 3, "end": 3.5, "intermediate": 1.5, "start": 1.5}.get(node.shape, 1)
        sp.line.width = Pt(width)
        if node.shape in ("doublecircle", "intermediate"):
            sp.line._get_or_add_ln().set("cmpd", "dbl")
        if _dashed(style):
            sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        color = parse_color(style.get("color", "")) or (_contrast(fill) if fill else pal["node_text"])
        if b.core:          # BPMN gateway/event: the text goes beside the symbol, on the -cross side
            if lay.lines[n]:
                gap = OUTSIDE_GAP * lay.font_size
                if d.direction in ("TB", "BT"):
                    lb, align, anchor = Box(b.x, b.y, (b.w - b.core) / 2 - gap, b.h), PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE
                else:
                    lb, align, anchor = Box(b.x, b.y, b.w, (b.h - b.core) / 2 - gap), PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM
                tb = shapes.add_textbox(X(lb.x), Y(lb.y), E(lb.w), E(lb.h))
                tb.name = f"text {n}"
                tf = tb.text_frame
                tf.word_wrap, tf.auto_size, tf.vertical_anchor = False, MSO_AUTO_SIZE.NONE, anchor
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                _write_text(tf, lay.lines[n], fs, parse_color(style.get("color", "")) or pal["text"], font,
                            align=align)
            continue
        tf = sp.text_frame
        tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Pt(0.2 * fs)
        tf.margin_top = tf.margin_bottom = Pt(0.1 * fs)
        _write_text(tf, lay.lines[n], fs, color, font)
        if node.icon:
            user_icon(b.x + 0.3 * lf, b.y + 0.3 * lf, icon, color, f"icon {n}")

    for r in lay.routes:
        e = r.edge
        if e.line == "invisible":
            continue
        pts = simplify([(X(x) / EMU_PER_PT, Y(y) / EMU_PER_PT) for x, y in r.points])
        color = parse_color(e.style.get("stroke", "")) or pal["edge"]
        width = _px_to_pt(e.style.get("stroke-width", "")) or (2.25 if e.line == "thick" else 1.0)
        if len(pts) - 1 <= 5:
            g = connector_geometry(pts)
            cx = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 1, 1)
            _strip_style(cx)
            _set_connector_xml(cx, g, (ids[e.src], r.src_site), (ids[e.dst], r.dst_site))
            line = cx.line
        else:       # more bends than a connector preset has: an unglued polyline
            fb = shapes.build_freeform(round(pts[0][0] * EMU_PER_PT), round(pts[0][1] * EMU_PER_PT), scale=1.0)
            fb.add_line_segments([(round(x * EMU_PER_PT), round(y * EMU_PER_PT)) for x, y in pts[1:]],
                                 close=False)
            cx = fb.convert_to_shape()
            _strip_style(cx)
            cx.fill.background()
            line = cx.line
        cx.name = f"edge {e.src}->{e.dst}"
        line.width = Pt(width)
        _set_color(line.color, color)
        if e.line == "dotted":
            line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ln = line._get_or_add_ln()
        for tag, marker in (("a:headEnd", e.start), ("a:tailEnd", e.end)):
            if marker != "none":
                etree.SubElement(ln, qn(tag), type=_MARKER_XML[marker], w="med", len="med")

    for r in lay.routes:
        if not r.label or r.edge.line == "invisible":
            continue
        b = r.label
        tb = shapes.add_textbox(X(b.x), Y(b.y), E(b.w), E(b.h))
        tb.name = f"label {r.edge.src}->{r.edge.dst}"
        _set_color_fill(tb, None, pal["label_fill"])
        tf = tb.text_frame
        tf.word_wrap, tf.auto_size, tf.vertical_anchor = False, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        color = parse_color(r.edge.style.get("color", "")) or pal["text"]
        _write_text(tf, r.label_lines, lfs, color, font)

    if grp is not None:     # child space == group space: no scaling, whatever python-pptx computed
        xfrm = grp._element.grpSpPr.find(qn("a:xfrm"))
        box = {"x": str(X(0)), "y": str(Y(0))}
        ext = {"cx": str(max(E(lay.width), 1)), "cy": str(max(E(lay.height), 1))}
        for tag, attrs in (("a:off", box), ("a:ext", ext), ("a:chOff", box), ("a:chExt", ext)):
            el = xfrm.find(qn(tag))
            for k, v in attrs.items():
                el.set(k, v)
    return prs


def _set_color_fill(shape, css: str | None, default):
    shape.fill.solid()
    _set_color(shape.fill.fore_color, (parse_color(css) if css else None) or default)


def _set_connector_xml(cx, g: dict, start: tuple[int, int | None], end: tuple[int, int | None]):
    el = cx._element
    sp_pr = el.find(qn("p:spPr"))
    xfrm = sp_pr.find(qn("a:xfrm"))
    for attr in ("rot", "flipH", "flipV"):
        xfrm.attrib.pop(attr, None)
    if g["rot"]:
        xfrm.set("rot", str(g["rot"]))
    if g["flipV"]:
        xfrm.set("flipV", "1")
    xfrm.find(qn("a:off")).set("x", str(round(g["x"] * EMU_PER_PT)))
    xfrm.find(qn("a:off")).set("y", str(round(g["y"] * EMU_PER_PT)))
    xfrm.find(qn("a:ext")).set("cx", str(round(g["w"] * EMU_PER_PT)))
    xfrm.find(qn("a:ext")).set("cy", str(round(g["h"] * EMU_PER_PT)))
    geom = sp_pr.find(qn("a:prstGeom"))
    geom.set("prst", g["prst"])
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for child in list(av):
        av.remove(child)
    for k, v in enumerate(g["adj"], 1):
        etree.SubElement(av, qn("a:gd"), name=f"adj{k}", fmla=f"val {v}")
    c_nv = el.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
    for tag, (shape_id, idx) in (("a:stCxn", start), ("a:endCxn", end)):
        if idx is not None:
            etree.SubElement(c_nv, qn(tag), id=str(shape_id), idx=str(idx))


# --------------------------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------------------------

def convert(src: str, opts: Options | None = None) -> tuple[Presentation, Diagram, Layout]:
    opts = opts or Options()
    d = parse(src)
    if opts.direction:
        d.direction = DIRECTIONS[opts.direction.upper()]
    if opts.render == "bpmn":
        to_bpmn(d, opts.events, opts.lanes)
    lay = layout(d, opts.font_size)
    return render(d, lay, opts), d, lay


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="mermaid2pptx", description=__doc__.split("\n\n")[0],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="Mermaid file (.mmd, or .md with ```mermaid blocks); '-' = stdin")
    ap.add_argument("-o", "--output", help="output .pptx (default: input name with .pptx)")
    ap.add_argument("--block", type=int, default=1, help="which ```mermaid block of a Markdown file (1-based)")
    ap.add_argument("--direction", choices=sorted(DIRECTIONS), help="override the diagram direction")
    ap.add_argument("--font-size", type=float, default=12.0, help="node text size in pt before fitting (12)")
    ap.add_argument("--font", help="font name (default: the theme font, so a paste adopts the target deck's)")
    ap.add_argument("--render", choices=("bpmn", "mermaid"), default="bpmn",
                    help="bpmn (default): tasks, gateways and events like a BPMN process; "
                         "mermaid: the flowchart's own shapes")
    ap.add_argument("--no-events", action="store_true",
                    help="bpmn: don't add the start and end events the diagram lacks")
    ap.add_argument("--no-lanes", action="store_true",
                    help="bpmn: draw top-level subgraphs as groups, never as swim lanes")
    ap.add_argument("--color", default="purple",
                    help="box colours: purple (default), slate, theme (the target deck's "
                         "theme colours) or #RRGGBB")
    ap.add_argument("--author", help=f"document author (default: ${AUTHOR_ENV}, else empty)")
    ap.add_argument("--aspect", choices=sorted(SLIDE_SIZES), default="16:9", help="slide shape (16:9)")
    ap.add_argument("--no-fit", action="store_true",
                    help="keep the natural size and size the slide to the diagram instead of shrinking it")
    ap.add_argument("--no-group", action="store_true", help="leave the shapes ungrouped on the slide")
    ap.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    args = ap.parse_args(argv)

    try:
        if args.input == "-":
            sys.stdin.reconfigure(encoding="utf-8-sig")      # the Windows console code page is not UTF-8
            src = sys.stdin.read()
        else:
            src = Path(args.input).read_text(encoding="utf-8-sig")   # tolerate Notepad's BOM
    except OSError as exc:
        print(f"error: cannot read {args.input}: {exc.strerror}", file=sys.stderr)
        return 1
    except UnicodeDecodeError:
        print(f"error: {args.input} is not UTF-8: save it as UTF-8 and try again", file=sys.stderr)
        return 1
    blocks = extract_blocks(src)
    if not 1 <= args.block <= len(blocks):
        ap.error(f"--block {args.block}: the input has {len(blocks)} mermaid block(s)")
    out = Path(args.output) if args.output else \
        (Path("diagram.pptx") if args.input == "-" else Path(args.input).with_suffix(".pptx"))
    try:
        palette(args.color)
    except ValueError as exc:
        ap.error(f"--color: {exc}")
    opts = Options(font_size=args.font_size, font=args.font, color=args.color, render=args.render,
                   events=not args.no_events, lanes=not args.no_lanes, author=args.author, group=not args.no_group,
                   fit=not args.no_fit, aspect=args.aspect, direction=args.direction)
    try:
        prs, d, lay = convert(blocks[args.block - 1], opts)
    except MermaidError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    try:
        prs.save(out)
    except OSError as exc:
        hint = " (is it open in PowerPoint?)" if isinstance(exc, PermissionError) else ""
        print(f"error: cannot write {out}: {exc.strerror}{hint}", file=sys.stderr)
        return 1
    for w in d.warnings:
        print(f"warning: {w}", file=sys.stderr)
    print(f"{out}: {len(d.nodes)} nodes, {len(d.edges)} edges, {len(d.subgraphs)} subgraphs")
    return 0


def _hold_window():
    """Keep the console open when the Windows executable was started from Explorer.

    A double-click, or a file dropped on the .exe, opens a console that closes as soon as the
    program ends, taking the message with it. From Explorer, the only processes attached to that
    console are the executable's own two (a one-file build runs under its bootloader); from a
    shell, the shell is attached as well.
    """
    if sys.platform != "win32" or not getattr(sys, "frozen", False):
        return
    import ctypes
    attached = ctypes.windll.kernel32.GetConsoleProcessList((ctypes.c_uint * 4)(), 4)
    if 0 < attached <= 2 and sys.stdin and sys.stdin.isatty():     # not in CI, not piped
        input("\nPress Enter to close this window.")


if __name__ == "__main__":
    try:
        sys.exit(main())
    finally:
        _hold_window()
