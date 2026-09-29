"""A laid-out sequence diagram -> shapes: participants, lifelines, messages, bars, notes, blocks."""
from __future__ import annotations

from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from ..geometry import Box
from ..options import Options, palette
from ..slide.connectors import add_path, simplify
from ..slide.deck import EMU_PER_PT, Canvas
from ..slide.shapes import frame, set_color, set_fill, strip_style, style_line, user_icon, write_text
from .layout import LABEL, SequenceLayout
from .model import Sequence

_RECT_TOP, _RECT_BOTTOM = 0, 2          # connection sites of a rectangle (and of a text box)


def render(s: Sequence, lay: SequenceLayout, opts: Options) -> Presentation:
    pal = palette(opts.color)
    cv = Canvas(opts, lay.width, lay.height)
    fs, lfs = cv.font(lay.font_size), cv.font(LABEL * lay.font_size)
    font = opts.font
    for b, lines, color in lay.groups:
        sp = _box(cv, MSO_SHAPE.RECTANGLE, b, "box", color or pal["cluster_fill"], pal["cluster_line"], 0.75)
        frame(sp.text_frame, MSO_ANCHOR.TOP, 0.2 * fs)
        write_text(sp.text_frame, lines, fs, pal["cluster_text"], font, bold=True)
    for b, color in lay.rects:
        _box(cv, MSO_SHAPE.RECTANGLE, b, "rect", color, None, 0)
    for blk in lay.blocks:
        _block_frame(cv, blk, pal)
    heads = _heads(cv, s, lay, pal, fs, font)
    for p, (lx, y0, y1) in lay.lifelines.items():
        cx = add_path(cv.shapes, _slide_points(cv, [(lx, y0), (lx, y1)]),
                      (heads[p][0], _RECT_BOTTOM), (heads[p][1], _RECT_TOP))
        cx.name = f"lifeline {p}"
        style_line(cx.line, 0.75, pal["cluster_line"], MSO_LINE_DASH_STYLE.DASH)
    for p, b in lay.activations:
        _box(cv, MSO_SHAPE.RECTANGLE, b, f"activation {p}", pal["cluster_fill"], pal["edge"], 0.75)
    for blk in lay.blocks:                  # over the bars, so they stay readable
        _block_labels(cv, blk, pal, lfs, font)
    for m in lay.messages:
        e = m.message
        cx = add_path(cv.shapes, _slide_points(cv, m.points), (0, None), (0, None))
        cx.name = f"message {e.src}->{e.dst}"
        style_line(cx.line, 1.0, pal["edge"], MSO_LINE_DASH_STYLE.DASH if e.line == "dotted" else None,
                   e.start, e.end)
        if m.label:
            tb = _text_box(cv, m.label, f"label {e.src}->{e.dst}", MSO_ANCHOR.BOTTOM)
            set_fill(tb, None, pal["label_fill"])         # lifelines and bars stop at the text
            align = PP_ALIGN.LEFT if e.src == e.dst else PP_ALIGN.CENTER
            write_text(tb.text_frame, m.lines, lfs, pal["text"], font, align=align)
        if m.number:
            sp = _box(cv, MSO_SHAPE.OVAL, m.number, f"number {m.index}", pal["edge"], None, 0)
            frame(sp.text_frame)
            write_text(sp.text_frame, [str(m.index)], max(5.0, lfs * 0.8), pal["label_fill"], font, bold=True)
    for b, lines in lay.notes:
        sp = _box(cv, MSO_SHAPE.FOLDED_CORNER, b, "note", pal["cluster_fill"], pal["cluster_line"], 0.75)
        frame(sp.text_frame, MSO_ANCHOR.MIDDLE, 0.2 * fs, wrap=True)
        write_text(sp.text_frame, lines, lfs, pal["cluster_text"], font)
    if lay.title:
        tb = _text_box(cv, lay.title[0], "title", MSO_ANCHOR.TOP)
        write_text(tb.text_frame, lay.title[1], cv.font(1.2 * lay.font_size), pal["text"], font, bold=True)
    return cv.finish()


def _slide_points(cv: Canvas, pts):
    return simplify([(cv.X(x) / EMU_PER_PT, cv.Y(y) / EMU_PER_PT) for x, y in pts])


def _box(cv: Canvas, preset, b: Box, name: str, fill, line, width: float):
    sp = cv.shapes.add_shape(preset, cv.X(b.x), cv.Y(b.y), cv.E(b.w), cv.E(b.h))
    strip_style(sp)
    sp.name = name
    set_fill(sp, None, fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.width = Pt(width)
        set_color(sp.line.color, line)
    return sp


def _text_box(cv: Canvas, b: Box, name: str, anchor):
    tb = cv.shapes.add_textbox(cv.X(b.x), cv.Y(b.y), cv.E(b.w), cv.E(b.h))
    tb.name = name
    frame(tb.text_frame, anchor)
    return tb


def _heads(cv: Canvas, s: Sequence, lay: SequenceLayout, pal: dict, fs: float, font) -> dict[str, tuple[int, int]]:
    """Draws each participant at the top and the bottom; returns their shape ids (top, bottom)."""
    ids = {}
    icon = 2.2 * lay.font_size
    for p, (top, bottom) in lay.heads.items():
        pair = []
        for b, where in ((top, "top"), (bottom, "bottom")):
            if s.participants[p].kind == "actor":       # a figure over the name, no box
                user_icon(cv, b.cx - icon / 2, b.y, icon, pal["node_fill"], f"actor icon {p}")
                sp = _text_box(cv, Box(b.x, b.y + icon, b.w, b.h - icon), f"actor {p} {where}", MSO_ANCHOR.MIDDLE)
                write_text(sp.text_frame, lay.lines[p], fs, pal["text"], font)
            else:
                sp = _box(cv, MSO_SHAPE.RECTANGLE, b, f"participant {p} {where}", pal["node_fill"], pal["node_line"], 1)
                frame(sp.text_frame, MSO_ANCHOR.MIDDLE, 0.2 * fs, wrap=True)
                write_text(sp.text_frame, lay.lines[p], fs, pal["node_text"], font)
            pair.append(sp.shape_id)
        ids[p] = tuple(pair)
    return ids


def _block_frame(cv: Canvas, blk, pal: dict):
    sp = cv.shapes.add_shape(MSO_SHAPE.RECTANGLE, cv.X(blk.box.x), cv.Y(blk.box.y), cv.E(blk.box.w), cv.E(blk.box.h))
    strip_style(sp)
    sp.name = f"{blk.kind} block"
    sp.fill.background()
    sp.line.width = Pt(0.75)
    set_color(sp.line.color, pal["edge"])
    for y, _, _ in blk.dividers:
        cx = add_path(cv.shapes, _slide_points(cv, [(blk.box.x, y), (blk.box.x + blk.box.w, y)]), (0, None), (0, None))
        cx.name = f"{blk.kind} divider"
        style_line(cx.line, 0.75, pal["edge"], MSO_LINE_DASH_STYLE.DASH)


def _block_labels(cv: Canvas, blk, pal: dict, lfs: float, font):
    tab = _box(cv, MSO_SHAPE.RECTANGLE, blk.tab, f"{blk.kind} tab", pal["cluster_fill"], pal["edge"], 0.75)
    frame(tab.text_frame)
    write_text(tab.text_frame, [blk.kind], lfs, pal["text"], font, bold=True)
    if blk.label:
        tb = _text_box(cv, blk.label, f"{blk.kind} condition", MSO_ANCHOR.MIDDLE)
        write_text(tb.text_frame, blk.lines, lfs, pal["text"], font, bold=True, align=PP_ALIGN.LEFT)
    for _, lines, b in blk.dividers:
        if b:
            tb = _text_box(cv, b, f"{blk.kind} condition", MSO_ANCHOR.MIDDLE)
            set_fill(tb, None, pal["label_fill"])
            write_text(tb.text_frame, lines, lfs, pal["text"], font, bold=True)
