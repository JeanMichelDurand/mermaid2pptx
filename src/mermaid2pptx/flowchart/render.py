"""A laid-out flowchart -> shapes on the slide: lanes, subgraphs, nodes, connectors, labels."""
from __future__ import annotations

from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from ..geometry import Box
from ..options import Options, palette
from ..slide.connectors import add_path, simplify
from ..slide.deck import EMU_PER_PT, Canvas
from ..slide.shapes import frame, set_color, set_fill, strip_style, style_line, user_icon, write_text
from ..styles import contrast, dashed, parse_color, px_to_pt
from ..text import text_width
from .layout import ICON_EM, OUTSIDE_GAP, Layout
from .model import Diagram
from .shapes import SHAPES


def render(d: Diagram, lay: Layout, opts: Options) -> Presentation:
    pal = palette(opts.color)
    cv = Canvas(opts, lay.width, lay.height)
    fs, lfs = cv.font(lay.font_size), cv.font(0.85 * lay.font_size)
    _lanes(cv, d, lay, pal, fs, opts.font)
    _clusters(cv, d, lay, pal, fs, opts)
    ids = _nodes(cv, d, lay, pal, fs, opts.font)
    _edges(cv, lay, pal, ids)
    _labels(cv, lay, pal, lfs, opts.font)
    return cv.finish()


def _lanes(cv: Canvas, d: Diagram, lay: Layout, pal: dict, fs: float, font: str | None):
    X, Y, E = cv.X, cv.Y, cv.E
    lf = lay.font_size
    icon = ICON_EM * lf
    vertical = d.direction in ("TB", "BT")
    for sid, body, header, title in lay.lanes:
        sg = d.subgraphs[sid]
        sp = cv.shapes.add_shape(MSO_SHAPE.RECTANGLE, X(body.x), Y(body.y), E(body.w), E(body.h))
        strip_style(sp)
        sp.name = f"lane {sid}"
        set_fill(sp, sg.style.get("fill"), pal["cluster_fill"])
        sp.line.width = Pt(0.75)
        set_color(sp.line.color, parse_color(sg.style.get("stroke", "")) or pal["cluster_line"])
        hd = cv.shapes.add_shape(MSO_SHAPE.RECTANGLE, X(header.x), Y(header.y), E(header.w), E(header.h))
        strip_style(hd)
        hd.name = f"lane header {sid}"
        set_fill(hd, None, pal["node_fill"])
        hd.line.width = Pt(0.75)
        set_color(hd.line.color, pal["node_line"])
        tf = hd.text_frame
        frame(tf, MSO_ANCHOR.MIDDLE, 0.2 * fs, wrap=True)
        if not vertical:
            tf._txBody.find(qn("a:bodyPr")).set("vert", "vert270")
        write_text(tf, title, fs, pal["node_text"], font, bold=True)
        if sg.icon:     # beside the title's first line, on its reading side
            tw = max(text_width(ln, lf) for ln in title)
            th = len(title) * lf * 1.2
            if vertical:
                user_icon(cv, header.cx - tw / 2 - icon - 0.3 * lf, header.cy - th / 2 + 0.1 * lf, icon,
                          pal["node_text"], f"icon {sid}")
            else:
                user_icon(cv, header.cx - th / 2 + 0.1 * lf, header.cy + tw / 2 + 0.3 * lf, icon,
                          pal["node_text"], f"icon {sid}")


def _clusters(cv: Canvas, d: Diagram, lay: Layout, pal: dict, fs: float, opts: Options):
    for sid, b, title in lay.clusters:
        sp = cv.shapes.add_shape(MSO_SHAPE.RECTANGLE, cv.X(b.x), cv.Y(b.y), cv.E(b.w), cv.E(b.h))
        strip_style(sp)
        sp.name = f"subgraph {sid}"
        style = d.subgraphs[sid].style
        set_fill(sp, style.get("fill"), pal["cluster_fill"])
        sp.line.width = Pt(px_to_pt(style.get("stroke-width", "")) or 0.75)
        set_color(sp.line.color, parse_color(style.get("stroke", "")) or pal["cluster_line"])
        if opts.render == "bpmn":       # a BPMN group: dash-dot border
            sp.line.dash_style = MSO_LINE_DASH_STYLE.LONG_DASH_DOT
        tf = sp.text_frame
        tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.TOP
        tf.margin_top = tf.margin_bottom = Pt(0.25 * fs)
        tf.margin_left = tf.margin_right = Pt(0.4 * fs)
        write_text(tf, title, fs, pal["cluster_text"], opts.font, bold=True,   # BPMN: label top-left
                   align=PP_ALIGN.LEFT if opts.render == "bpmn" else PP_ALIGN.CENTER)


def _nodes(cv: Canvas, d: Diagram, lay: Layout, pal: dict, fs: float, font: str | None) -> dict[str, int]:
    """Draws the nodes; returns their shape ids, for gluing the connectors."""
    X, Y, E = cv.X, cv.Y, cv.E
    lf = lay.font_size
    ids: dict[str, int] = {}
    for n, b in lay.boxes.items():
        node = d.nodes[n]
        preset, _, flip_h = SHAPES[node.shape]
        shape_box = Box(b.cx - b.core / 2, b.cy - b.core / 2, b.core, b.core) if b.core else b
        sp = cv.shapes.add_shape(preset, X(shape_box.x), Y(shape_box.y), E(shape_box.w), E(shape_box.h))
        strip_style(sp)
        sp.name = f"node {n}"
        ids[n] = sp.shape_id
        if flip_h:
            sp._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
        if node.shape in ("task", "call"):
            sp.adjustments[0] = 0.12
        style = d.node_style(node)
        event = node.shape in ("start", "end", "intermediate")
        fill = parse_color(style.get("fill", ""))
        set_fill(sp, style.get("fill"), pal["label_fill"] if event else pal["node_fill"])
        stroke = parse_color(style.get("stroke", ""))
        set_color(sp.line.color, stroke or (pal["accent"] if event else pal["node_line"]))
        width = px_to_pt(style.get("stroke-width", "")) or \
            {"doublecircle": 3, "call": 3, "end": 3.5, "intermediate": 1.5, "start": 1.5}.get(node.shape, 1)
        sp.line.width = Pt(width)
        if node.shape in ("doublecircle", "intermediate"):
            sp.line._get_or_add_ln().set("cmpd", "dbl")
        if dashed(style):
            sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        color = parse_color(style.get("color", "")) or (contrast(fill) if fill else pal["node_text"])
        if b.core:          # BPMN gateway/event: the text goes beside the symbol, on the -cross side
            if lay.lines[n]:
                gap = OUTSIDE_GAP * lf
                if d.direction in ("TB", "BT"):
                    lb, align, anchor = Box(b.x, b.y, (b.w - b.core) / 2 - gap, b.h), PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE
                else:
                    lb, align, anchor = Box(b.x, b.y, b.w, (b.h - b.core) / 2 - gap), PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM
                tb = cv.shapes.add_textbox(X(lb.x), Y(lb.y), E(lb.w), E(lb.h))
                tb.name = f"text {n}"
                frame(tb.text_frame, anchor)
                write_text(tb.text_frame, lay.lines[n], fs, parse_color(style.get("color", "")) or pal["text"],
                           font, align=align)
            continue
        tf = sp.text_frame
        tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Pt(0.2 * fs)
        tf.margin_top = tf.margin_bottom = Pt(0.1 * fs)
        write_text(tf, lay.lines[n], fs, color, font)
        if node.icon:
            user_icon(cv, b.x + 0.3 * lf, b.y + 0.3 * lf, ICON_EM * lf, color, f"icon {n}")
    return ids


def _edges(cv: Canvas, lay: Layout, pal: dict, ids: dict[str, int]):
    for r in lay.routes:
        e = r.edge
        if e.line == "invisible":
            continue
        pts = simplify([(cv.X(x) / EMU_PER_PT, cv.Y(y) / EMU_PER_PT) for x, y in r.points])
        cx = add_path(cv.shapes, pts, (ids[e.src], r.src_site), (ids[e.dst], r.dst_site))
        cx.name = f"edge {e.src}->{e.dst}"
        style_line(cx.line, px_to_pt(e.style.get("stroke-width", "")) or (2.25 if e.line == "thick" else 1.0),
                   parse_color(e.style.get("stroke", "")) or pal["edge"],
                   MSO_LINE_DASH_STYLE.DASH if e.line == "dotted" else None, e.start, e.end)


def _labels(cv: Canvas, lay: Layout, pal: dict, lfs: float, font: str | None):
    for r in lay.routes:
        if not r.label or r.edge.line == "invisible":
            continue
        b = r.label
        tb = cv.shapes.add_textbox(cv.X(b.x), cv.Y(b.y), cv.E(b.w), cv.E(b.h))
        tb.name = f"label {r.edge.src}->{r.edge.dst}"
        set_fill(tb, None, pal["label_fill"])
        frame(tb.text_frame)
        write_text(tb.text_frame, r.label_lines, lfs, parse_color(r.edge.style.get("color", "")) or pal["text"], font)
