"""A laid-out Gantt chart -> shapes: section bands, grid, bars, milestones, axis labels."""
from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from ..geometry import Box
from ..options import Options, palette
from ..slide.connectors import add_path
from ..slide.deck import EMU_PER_PT, Canvas
from ..slide.shapes import frame, set_color, set_fill, strip_style, style_line, write_text
from .layout import LABEL, GanttLayout
from .model import Gantt


def render(g: Gantt, lay: GanttLayout, opts: Options) -> Presentation:
    pal = palette(opts.color)
    cv = Canvas(opts, lay.width, lay.height)
    fs, lfs = cv.font(lay.font_size), cv.font(LABEL * lay.font_size)
    font = opts.font
    for k, (name, band, col, lines) in enumerate(lay.sections):
        fill = pal["cluster_fill"] if k % 2 == 0 else pal["label_fill"]       # alternate bands
        _shape(cv, MSO_SHAPE.RECTANGLE, band, f"section {name}", fill, None)
        if lines:
            tb = _text(cv, col, f"section label {name}", MSO_ANCHOR.MIDDLE)
            tb.text_frame.margin_left = Pt(0.4 * fs)
            write_text(tb.text_frame, lines, fs, pal["cluster_text"], font, bold=True, align=PP_ALIGN.LEFT)
    for b in lay.shaded:
        sp = _shape(cv, MSO_SHAPE.RECTANGLE, b, "excluded days", pal["cluster_line"], None)
        _transparency(sp, 70)
    top, bottom = lay.grid
    for x, label in lay.ticks:
        pts = [(cv.X(x) / EMU_PER_PT, cv.Y(top) / EMU_PER_PT), (cv.X(x) / EMU_PER_PT, cv.Y(bottom) / EMU_PER_PT)]
        cx = add_path(cv.shapes, pts, (0, None), (0, None))
        cx.name = f"grid {label}"
        style_line(cx.line, 0.5, pal["cluster_line"], MSO_LINE_DASH_STYLE.DASH)
        tb = _text(cv, Box(x - lay.tick_label_w / 2, lay.axis_y + 0.2 * lay.font_size, lay.tick_label_w,
                           1.4 * LABEL * lay.font_size), f"tick {label}", MSO_ANCHOR.TOP)
        write_text(tb.text_frame, [label], lfs, pal["text"], font)
    for b in lay.bars:
        t = b.task
        fill, text, line = pal["node_fill"], pal["node_text"], pal["node_line"]
        if "done" in t.tags:
            fill, text, line = pal["done_fill"], pal["cluster_text"], pal["cluster_line"]
        elif "active" in t.tags:
            fill, text, line = pal["label_fill"], pal["cluster_text"], pal["node_fill"]
        if "crit" in t.tags:
            line = pal["crit"]
            if "done" not in t.tags and "active" not in t.tags:
                fill, text = pal["crit"], pal["crit_text"]
        preset = MSO_SHAPE.DIAMOND if t.milestone else MSO_SHAPE.ROUNDED_RECTANGLE
        sp = _shape(cv, preset, b.box, f"{'milestone' if t.milestone else 'task'} {t.id}", fill, line,
                    2.0 if {"crit", "active"} & t.tags else 1.0)
        if not t.milestone:
            sp.adjustments[0] = 0.15
        if b.inside:
            frame(sp.text_frame)
            write_text(sp.text_frame, [t.name], lfs, text, font)
        else:
            tb = _text(cv, b.label, f"label {t.id}", MSO_ANCHOR.MIDDLE)
            left = b.label.x >= b.box.x + b.box.w - 0.01
            write_text(tb.text_frame, [t.name], lfs, pal["text"], font, align=PP_ALIGN.LEFT if left else PP_ALIGN.RIGHT)
    if lay.title:
        tb = _text(cv, lay.title[0], "title", MSO_ANCHOR.TOP)
        write_text(tb.text_frame, lay.title[1], cv.font(1.2 * lay.font_size), pal["text"], font, bold=True)
    return cv.finish()


def _shape(cv: Canvas, preset, b: Box, name: str, fill, line, width: float = 1.0):
    sp = cv.shapes.add_shape(preset, cv.X(b.x), cv.Y(b.y), cv.E(b.w), cv.E(b.h))
    strip_style(sp)
    sp.name = name
    set_fill(sp, None, fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.width = Pt(width)
        set_color(sp.line.color, line)
    return sp


def _text(cv: Canvas, b: Box, name: str, anchor):
    tb = cv.shapes.add_textbox(cv.X(b.x), cv.Y(b.y), cv.E(b.w), cv.E(b.h))
    tb.name = name
    frame(tb.text_frame, anchor)
    return tb


def _transparency(shape, percent: int):
    """Lets what is behind show through the solid fill."""
    clr = shape.fill._xPr.find(qn("a:solidFill"))[0]
    etree.SubElement(clr, qn("a:alpha"), val=str((100 - percent) * 1000))
