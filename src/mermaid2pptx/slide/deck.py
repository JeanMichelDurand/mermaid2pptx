"""The deck: one blank slide, the diagram scaled to fit it, in one group."""
from __future__ import annotations

import io
import os
import zipfile
from datetime import datetime, timezone

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from ..errors import MermaidError
from ..options import AUTHOR_ENV, MARGIN, SLIDE_SIZES, TEMPLATE_ENV, Options

EMU_PER_PT = 12700
_POTX = b"presentationml.template.main+xml"
_PPTX = b"presentationml.presentation.main+xml"


def open_template(path: str | None) -> Presentation:
    """The template deck emptied of its slides, or python-pptx's blank deck when `path` is empty.

    A .potx is read as a .pptx: python-pptx refuses its main part's content type, the only difference.
    """
    if not path:
        return Presentation()
    try:
        with zipfile.ZipFile(path) as src:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
                for item in src.infolist():
                    data = src.read(item)
                    dst.writestr(item, data.replace(_POTX, _PPTX) if item.filename == "[Content_Types].xml" else data)
        prs = Presentation(buf)
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise MermaidError(f"template {path}: not a PowerPoint .pptx or .potx file ({exc})") from exc
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        ids.remove(sld)
    return prs


def blank_layout(prs: Presentation):
    """The layout that puts the fewest placeholders on a slide: "Blank" in most templates."""
    return min(prs.slide_layouts, key=lambda layout: len(list(layout.iter_cloneable_placeholders())))


class Canvas:
    """A diagram of `width` x `height` points on the slide of a new deck.

    The deck is `opts.template` (or $MERMAID2PPTX_TEMPLATE) emptied of its slides, so the theme
    colours and fonts are the template's, else python-pptx's built-in blank presentation. Layouts
    work in points from (0, 0); X, Y and E turn a layout's x, y and length into slide EMU,
    scaled down to fit the slide unless `opts.fit` is off (then the slide fits the diagram).
    """

    def __init__(self, opts: Options, width: float, height: float, name: str = "Mermaid diagram"):
        self.width, self.height = width, height
        template = opts.template if opts.template is not None else os.environ.get(TEMPLATE_ENV, "")
        self.prs = prs = open_template(template)
        now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
        cp = prs.core_properties         # replace the built-in deck's own metadata
        cp.author = cp.last_modified_by = opts.author if opts.author is not None else os.environ.get(AUTHOR_ENV, "")
        cp.title, cp.subject, cp.comments, cp.keywords, cp.category = "", "", "", "", ""
        cp.created = cp.modified = now
        cp.revision = 1
        sw, sh = (prs.slide_width / EMU_PER_PT, prs.slide_height / EMU_PER_PT) if template \
            else SLIDE_SIZES[opts.aspect]
        if opts.fit:
            s = min(1.0, (sw - 2 * MARGIN) / max(width, 1), (sh - 2 * MARGIN) / max(height, 1))
        else:
            s = 1.0
            sw, sh = max(width + 2 * MARGIN, 72), max(height + 2 * MARGIN, 72)
        prs.slide_width, prs.slide_height = Emu(round(sw * EMU_PER_PT)), Emu(round(sh * EMU_PER_PT))
        self.s = s
        self.ox, self.oy = (sw - width * s) / 2, (sh - height * s) / 2
        slide = prs.slides.add_slide(blank_layout(prs))
        for ph in list(slide.placeholders):          # a template may have no layout without any
            ph._element.getparent().remove(ph._element)
        self.group = slide.shapes.add_group_shape() if opts.group else None
        self.shapes = self.group.shapes if self.group is not None else slide.shapes
        if self.group is not None:
            self.group.name = name

    def E(self, v: float) -> int:
        """A length."""
        return int(round(v * self.s * EMU_PER_PT))

    def X(self, v: float) -> int:
        return int(round((self.ox + v * self.s) * EMU_PER_PT))

    def Y(self, v: float) -> int:
        return int(round((self.oy + v * self.s) * EMU_PER_PT))

    def font(self, size: float) -> float:
        """A font size once scaled, to the half point, never below 5 pt."""
        return max(5.0, round(size * self.s * 2) / 2)

    def finish(self) -> Presentation:
        if self.group is not None:     # child space == group space: no scaling, whatever python-pptx computed
            xfrm = self.group._element.grpSpPr.find(qn("a:xfrm"))
            box = {"x": str(self.X(0)), "y": str(self.Y(0))}
            ext = {"cx": str(max(self.E(self.width), 1)), "cy": str(max(self.E(self.height), 1))}
            for tag, attrs in (("a:off", box), ("a:ext", ext), ("a:chOff", box), ("a:chExt", ext)):
                el = xfrm.find(qn(tag))
                for k, v in attrs.items():
                    el.set(k, v)
        return self.prs
